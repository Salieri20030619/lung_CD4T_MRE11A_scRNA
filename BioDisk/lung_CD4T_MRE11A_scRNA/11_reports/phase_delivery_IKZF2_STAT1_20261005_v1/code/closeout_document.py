from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,hashlib
B=Path(__file__).parent
server_layout=B.name=='code' and (B.parent/'figures').exists()
D=B.parent/'deliverables' if server_layout else B/'closeout_delivery';D.mkdir(exist_ok=True)
figure1=B.parent/'figures/01_independent_TF_evidence.png' if server_layout else B/'ic_delivery/01_independent_TF_evidence.png'
figure2=B.parent/'figures/IKZF2_IL7R_fixed_deletion_summary.png' if server_layout else B/'ikdel_delivery/IKZF2_IL7R_fixed_deletion_summary.png'
doc=Document();s=doc.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21);s.top_margin=Cm(1.65);s.bottom_margin=Cm(1.55);s.left_margin=Cm(1.7);s.right_margin=Cm(1.7)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Microsoft YaHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 st.font.size=Pt(10 if name in ['Normal','Body Text'] else 10)
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(26);doc.styles['Title'].font.color.rgb=RGBColor.from_string('143C55')
doc.styles['Heading 1'].font.size=Pt(19);doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('143C55')
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.color.rgb=RGBColor.from_string('177E89')
doc.styles['Caption'].font.size=Pt(8);doc.styles['Caption'].font.color.rgb=RGBColor.from_string('52616B')
hp=s.header.paragraphs[0];hp.text='肺部疾病单细胞研究  /  人类优先阶段交付';hp.style='Caption'
fp=s.footer.paragraphs[0];fp.alignment=2;r=fp.add_run('2026-10-05  ·  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,2)
def page(t):doc.add_page_break();doc.add_heading(t,1)
def tab(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.autofit=False
 for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');pr.append(el)
  for i,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[i])
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.line_spacing=1.0
    for rr in pp.runs:rr.font.size=Pt(8.5)
 header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
 return t
doc.add_heading('人类 CD4 T 细胞变化与\n独立候选证据',0)
p('IKZF2 · STAT1  /  巨噬细胞通讯背景与补充候选', 'Subtitle')
p('阶段交付  |  2026年10月5日  |  基于既有结果收束')
h('本轮判断')
p('IKZF2 和 STAT1 均可保留为后续讨论的独立候选。IKZF2 的主要人类证据是特定亚群中来源校正后的 RNA 差异与程序关联；STAT1 的主要人类证据是调控子活性与炎症小体相关转录程序的关联，并有小鼠 RNA 差异作为跨物种背景。两者不需要组成同一条级联通路。')
tab(['候选','可保留的积极证据','结论边界'],[
 ['IKZF2','人 GZMA RNA：Source P=0.0021，FDR=0.0253；人 IL7R 活性–程序关联三种子均名义显著。','RNA 基础模型未达名义；IL7R 关联原结果未计算 FDR。'],
 ['STAT1','人 GZMA 活性–炎症小体相关程序：三种子 P=0.0086–0.0309；鼠 Core、ActCD4 RNA 名义上调。','上述人程序关联和鼠 RNA 的原 FDR 未通过；人疾病 RNA 未形成支持。']],[2,8,7.3])
h('从患者表型出发')
p('既有数据支持部分 CD4 组成和亚群内部转录程序的疾病相关变化，但来源与疾病高度混杂。报告先交代细胞范围、分母和供者，再分别呈现两位候选；巨噬细胞通讯及新独立候选作为背景或补充，不作为两位 TF 的独立验证。')
h('本次交付范围')
p('复用原始分析结果，保留名义支持、FDR 状态和未支持结果。此次未重跑 SCENIC、Bootstrap、NicheNet 或候选筛选；固定靶集删除仅作为已有计算敏感性分析。实验选择与目的基因机制阶段留待讨论。')
p('阅读提示：P 为原模型名义检验；FDR 为原检验家族的 BH 校正。未计算 FDR 与 FDR 未通过严格区分。本文所有 n 均指供者或独立文库，不将细胞、种子视为独立重复。','Caption')
page('01  患者背景与 CD4 表型')
p('人类 RNA 总名单为 19 位供者：10 Control、9 CWP。来源 G135 为 2 对照，G227 为 8 对照及 2 CWP，G288 为 7 CWP。CWP 保留原临床标签，不能直接等同严格诊断的矽肺。固定 AUC 17 名单与 RNA 19 名单不同；各亚群按既定细胞覆盖使用实际可用供者。')
tab(['范围','细胞数','≥20细胞供者（对照/CWP）','定义'],[['CCR7',1172,'9（6/3）','CD4-CCR7'],['GZMA',2554,'12（7/5）','CD4-GZMA'],['IL7R',1379,'16（7/9）','CD4-IL7R'],['Treg',600,'10（5/5）','原 Treg 注释'],['CD4state',818,'7（1/6）','疾病比较对照不足']],[2.6,2,5.5,7.2])
p('Primary 包含 CD4state；Core 排除 CD4state。组成比例以相应范围内 CD4 总细胞数为分母，不能把 Primary 和 Core 的比例混用。注释复用已核验对象；现有“免疫图景”是分析范围说明，不代表所有免疫细胞均已完成新的疾病检验。')
h('组成变化：基础模型的线索')
tab(['亚群 / 分母','CWP−Control 比例差','P / FDR','解释'],[['CCR7 / Primary','−0.1133','0.0303 / 0.0538','基础模型名义减少'],['Treg / Primary','+0.0752','0.0323 / 0.0538','基础模型名义增加'],['CD4state / Primary','+0.2174','0.0239 / 0.0538','基础模型名义增加']],[4,4,4,5.3])
p('以上均为 19 供者基础模型；来源校正后未保留名义支持，因此不能称为来源独立的稳定疾病组成特征。WithinG227 仅 2 病例，无法依靠不断增加协变量消除设计局限。')
h('亚群内部程序：保留亚群特异性')
p('GZMA 的 TReg 相关转录程序在基础模型中升高（n=12，效应 +0.0476，P=0.000425）；来源校正后仍为名义支持（效应 +0.0388，P=0.0131，原范围内12程序 FDR=0.1574）。CCR7 的 Inhib 程序在 SourceDep 中为名义支持（n=9，效应 +0.0524，P=0.0157，FDR=0.1884）。它们均不是完整 Treg 抑制功能或细胞功能实验结论。')
p('近期 Core/IL7R 五小面板的 30 项疾病比较均无名义显著；其配体相关结果不能因此被称为已证实疾病表型。来源：[S1]。','Caption')
page('02  IKZF2：两类人类独立关联')
h('GZMA 亚群 RNA：来源校正后出现差异')
p('供者为独立单位，n=12（5 CWP、7 对照）。结局是供者平均 log1p CP10k；下表效应是 OLS 调整后均值差，不是 edgeR log2FC。')
tab(['模型','效应 [95% CI]','P','FDR'],[['Base','+0.0108 [−0.0193, +0.0409]','0.4442','0.9603'],['Source','+0.0445 [+0.0209, +0.0682]','0.002109','0.025304'],['SourceDep','+0.0441 [+0.0183, +0.0699]','0.004281','0.051376']],[3,7.1,3.6,3.6])
p('Source 在该原家族中通过 FDR；SourceDep 为名义支持、FDR 略高于 0.05。基础模型未达名义，表明结论依赖来源调整，不能描述为跨来源一致的 RNA 上调。GZMA 中 IKZF2 RNA 仅在少量细胞检出（既有审计 29/2554），需保留稀疏检出与来源分布限制。')
h('IL7R 亚群活性：与固定背景扣除后的程序负相关')
p('n=16（9 CWP、7 对照）；结局 excess 为冻结的疾病上调程序分数减去三组固定匹配背景的平均分数。偏秩相关调整组别、来源和接收端深度。AUC 是调控子活性读数，不等于 IKZF2 RNA。')
tab(['DropOverlap 种子','偏相关 ρ','P','FDR'],[['17','−0.6349','0.026557','未计算'],['29','−0.7618','0.003981','未计算'],['43','−0.5977','0.040139','未计算']],[5,4,4,4.3])
p('三种子均为名义负相关；Original 三种子方向一致、P=0.0023–0.0292。种子是算法敏感性，不是独立患者复制；该结果与 GZMA RNA 属于不同亚群、不同检验，不应拼接成同一因果链。')
h('已有固定靶集删除：作为敏感性保留')
p('高权重 10%/20% 靶集删除后，六规格的剩余关联仍为名义支持（ρ=−0.622 至 −0.802，P=0.0017–0.0309）。但关联变化 Δρ 为 −0.088 至 +0.013，配对 jackknife 区间均含 0，未出现一致的高权重靶集特异依赖。该计算不是基因敲除，也不能证明疾病逆转。')
p('来源：[S2] 人类 RNA；[S3] IL7R 固定背景与删除敏感性。','Caption')
page('03  STAT1：程序关联与小鼠对照')
h('人 GZMA：炎症小体相关转录程序')
p('既有 DropOverlap 调控子活性与 score_Inflam 呈正相关。该模块包括 NLRP3、AIM2、NLRC4、PYCARD、CASP1；反映转录程序，不代表炎症小体组装、蛋白活化或焦亡功能。n=12（5 CWP、7 对照），SourceDep 调整组别、来源及接收端深度。')
tab(['种子','偏相关 ρ','P','原 FDR'],[['17','+0.8068','0.008587','0.256882'],['29','+0.7318','0.025022','0.344046'],['43','+0.7136','0.030860','0.309623']],[4,4.5,4.4,4.4])
p('三个种子均为名义支持，原 FDR 均未通过。Core 范围仅种子17在同一 SourceDep 模型名义支持（ρ=0.7313，P=0.0045，FDR=0.0607）；种子29/43的 P=0.1550/0.1805，提示范围和种子依赖。')
p('人 STAT1 RNA 在本次复核的 All 基础、来源及来源+深度模型中均未达名义疾病差异（可估计范围：CCR7、GZMA、IL7R、Treg、Core、Primary）。因此保留“活性–程序关联”，不写成“人 STAT1 疾病 RNA 上调通过”。')
h('鼠 GSE314943：3+3 独立文库的 RNA 差异')
tab(['范围 / 基因','edgeR log2FC','P','原 FDR'],[['Core / Ikzf2','+1.8279','0.001079','0.056631'],['Core / Stat1','+1.0632','0.001945','0.074361'],['ActCD4 / Ikzf2','−0.2068','0.778127','0.970349'],['ActCD4 / Stat1','+1.0424','0.003723','0.417933'],['Naive / Stat1','+0.8255','0.020267','0.952086']],[5,4.2,4,4.1])
p('以上均为 All 臂、edgeR QL 组别模型。鼠 Core 中两基因名义上调；更明确的 ActCD4 中仅 Stat1 保留名义支持，不能用 Core 的混合信号替代 ActCD4。各行均未通过原全转录组 FDR。鼠与人基因符号按既有映射对应，亚群并非自动等价。')
p('GSE183682 的 7/56 天每组仅 1 文库，只作方向参考，不提供正式组间显著性验证。来源：[S4]。','Caption')
page('04  巨噬细胞通讯：背景证据')
p('发送端沿用已审阅巨噬细胞范围：MacAM 为原簇0/2/11/13，MonoMac 为簇3/7，MacWide 为二者并集。MonoMac 保留混合单核/巨噬特征；MacWide 不是额外独立重复。T/NK 仅用于来源竞争敏感性，不替代巨噬发送端。')
tab(['发送范围','覆盖与注释限制','报告用途'],[['MacAM','9771细胞：CWP 687、对照9084；≥20细胞供者14（5/9）。簇2/11分别由单一对照供者约98.8%/97.2%主导。','患者巨噬发送背景；不能称疾病新亚型'],['MonoMac','簇3/7同时含 C1Q/MARCO/SPP1 与 FCN1 相关特征。','混合髓系候选，保留来源身份限制'],['MacWide','MacAM 与 MonoMac 的 union。','发送范围敏感性']],[3,9,5.3])
h('已有可保留的通信线索')
tab(['信号','已有证据','不能据此声称'],[['IL15','独立患者 AM/MLM 定位；鼠 Core FDR-Up 靶集 NicheNet 排名3。人 MonoMac→Treg 名义靶集的既有同供者关联 P=0.0327、FDR=0.4974。','IL15RA 转呈递已获验证；或已连通 IKZF2/STAT1'],['TNF','MonoMac→IL7R Activation 的既有同供者关联：组别+来源调整 ρ=0.5511，P=0.0269、FDR=0.2345。','对所有更强调整稳定，或已证明疾病功能'],['IL18','独立两患者数据支持 AM/MLM 发送端定位；接收端复合物 RNA 共检出有限。','低检出等于受体生物学不存在，或因预测获得因果']],[2.2,10.1,5])
p('这些条目是已有通讯背景摘要，不是新候选优选或两位 TF 的独立验证。完整同供者 n、模型、效应、P/FDR、靶集与背景保留在服务器原表；NicheNet 排名不附会为显著性 P 值。')
h('把四类证据分开')
p('① 配体/受体 RNA 与复合物兼容性；② NicheNet 配体活性与预测靶；③ 预测靶与真实疾病程序的吻合；④ 同供者关联。它们分别回答可表达、可预测、可对应和可关联的问题，不自动形成因果通路。')
p('沿用正式 nichenetr 2.2.1.1 配套 classic human prior（Zenodo 3260758），不改称 2022 human v2。原 10% 表达门槛、5% 敏感性、两端≥20细胞、至少2供者支持及名义/FDR靶集均保留。鼠 Core 主靶集与 ActCD4 名义探索靶集分开，ActCD4 FDR集合不足仍保留不足状态。来源：[S5]。','Caption')
page('05  补充：独立候选及其限制')
p('这些结果来自此前已完成的较宽候选筛选，本轮仅整理，不再扩筛。它们补充巨噬细胞与 CD4 程序的关联图景，不作为 IKZF2 或 STAT1 的验证。下表保留发现阶段的完整检验家族 FDR，不针对三条结果重新缩小家族。')
tab(['MacAM 候选 / 程序','供者（CWP/对照）','Source ρ','P / FDR'],[['HBEGF / IL7R Mitophagy','11（5/6）','+0.8477','0.00786 / 0.4069'],['ALCAM / IL7R Mitophagy','11（5/6）','−0.7951','0.01833 / 0.5974'],['HMGB1 / GZMA TReg','9（3/6）','+0.7922','0.03368 / 0.7655']],[6.1,4.5,2.4,4.3])
h('分解评分后，保留更准确的解释')
p('HBEGF 与 IL7R 未扣背景的原靶基因均值关联，在来源模型中 P=0.1519；加入两端深度后 P=0.0627。ALCAM 与同一原靶基因均值的来源模型 P=0.9203，提示其原信号明显依赖背景扣除。两者不能直接被称为促进或抑制线粒体自噬。')
p('HMGB1 与 GZMA 原 TReg 靶基因均值的来源模型 P=0.8815，GZMA 原靶基因均值本身的疾病比较 P=0.4656。相对程序的相关不等于真实 Treg 功能差异，更不等于免疫抑制机制。')
h('其他审计结果保留在补充层')
p('SERPINA1 在 MonoMac 混合髓系范围有较强发现期关联；分解到纯 MacAM 后不支持。因此保留混合髓系假说，不升级为纯巨噬细胞信号。TXLNA 的 n=6 完美相关原近似 P 值精度不可靠；已有精确置换修订和原值均保留，修订后不通过全家族 FDR。')
h('外部人类数据的用途')
p('GSE174725 为 2 Exposed + 3 Silicosis 的小队列，已有注释复核和固定验证结果，不构成 FDR 层面的独立确认。GSE296682 的两位患者无对照，可支持发送端定位，不能检验疾病差异。两者注释范围不同，不能直接互换巨噬/单核标签。')
p('完整补充报告与服务器分析包已交付并保留；本次不复制大规模候选表到正文。来源：[S6]。','Caption')
page('图示  两候选的独立证据与敏感性')
doc.add_picture(str(figure1),width=Cm(17.2))
p('图1｜A：人 GZMA 中 IKZF2 的 RNA OLS 疾病效应，线段为95%置信区间。B：人 IL7R 中 IKZF2 AUC 与固定匹配背景扣除后的 excess 程序的偏秩相关；O=Original，D=DropOverlap，s17/29/43=种子，线段为近似 Fisher 区间。C：鼠 Core 中 Ikzf2/Stat1 的 edgeR RNA log2FC。三个面板的结局、单位和模型不同，不可合并效应。q 为相应原FDR；IL7R原结果未计算FDR。来源：[S2–S4]。','Caption')
doc.add_picture(str(figure2),width=Cm(16.9))
p('图2｜IKZF2–IL7R 固定靶集删除敏感性。高权重10%/20%删除与每种子20个匹配删除对照沿用已有规格；种子是算法敏感性，非独立复制。删除后的剩余关联仍存在，但高权重删除引起的相关变化未显示一致特异依赖，配对jackknife区间均含0。target=靶基因；matched=匹配对照；Δρ=相对基线的相关变化。该图不是生物学KO。来源：[S3]。','Caption')
page('06  证据边界与后续讨论')
h('当前支持的内容')
p('人类 CD4 中存在亚群特异的组成及转录程序线索；IKZF2 与 STAT1 各自具有可报告的独立关联。IKZF2 的 GZMA RNA 来源模型在原家族中通过 FDR，其 IL7R 程序关联保留为名义支持。STAT1 的人 GZMA 程序关联与鼠 RNA 差异保留为名义支持。巨噬发送端定位、通讯预测和配对关联为候选背景。')
h('限制集中说明')
tab(['限制','对结论的实际影响'],[['来源与疾病混杂','G288无对照，G135无病例；WithinG227仅2病例。来源校正不能替代平衡设计或独立复制。'],['读数不同','RNA、AUCell、模块分数、比例、NicheNet排名分别解释；OLS效应不与edgeR log2FC拼接。'],['样本与覆盖','各亚群实际n不同；细胞不足/模型不可识别保留NA原因。受体低RNA检出不等于生物学不存在。'],['多重检验与敏感性','名义P<0.05不是全家族确认；三种子不是独立重复。未计算FDR不写成FDR阴性。'],['功能与因果','TReg、炎症小体、线粒体/焦亡评分是转录程序。固定靶集删除不是KO；配体预测和相关不证明上下游因果。'],['物种与疾病边界','CWP不等同严格矽肺；鼠Core混合范围不等于人单一亚群。1文库/组的时间数据只作方向。']],[4,13.3])
h('目的基因阶段：留待选择')
p('若后续进入机制验证，可分别选择 IKZF2 的 GZMA RNA 问题、IL7R 固定程序问题，或 STAT1 的特定 CD4 炎症小体相关程序问题，并先固定细胞范围、实验终点与必要对照。无需预设 IKZF2→STAT1 或配体→TF 级联。实验由客户根据样本和实验条件决定。')
p('本轮至此收束：不追加独立候选、不启动新的目的基因调控子或机制整合。原完整正负结果、成员表及复现材料继续保留在服务器。')
page('附录  模型与来源索引')
p('项目根目录 P = /mnt/BioDisk/wangmingyuan/lung_CD4T_MRE11A_scRNA。下列均为既有分析文件；报告汇总不重拟合模型。详细供者名单、原矩阵和完整 ZIP 不上传 Library。','Caption')
tab(['索引','模型 / FDR 规则','服务器定位（相对 P/11_reports）'],[
 ['S1','组成：固定分母内跨亚群BH；程序按原scope/roster/model家族。扩展程序池与旧池FDR不同，不混用。','human_first_20261004_v1/results/\nstage1_all_finite_models.tsv'],
 ['S2','人RNA OLS：Base=组别；Source加来源；SourceDep加深度。BH按readout/variant/seed/filter/model跨scope与TF；GZMA所示家族m=12。','human_alignment_20261005_v1/results/\nTF_disease_models.tsv'],
 ['S3','IL7R偏秩相关调整组别/来源/接收深度；excess原结果未计算FDR。删除剩余关联的BH不用于检验关联变化。','human_gap_diagnostics_20261005_v1/results/\nIKZF2_matched_reference_contrast_models.tsv；\nIKZF2_IL7R_target_deletion_20261005_v1/'],
 ['S4','旧TF程序关联BH在species/filter/seed/variant/scope/model内跨TF–feature对，沿用原n≥10规则；RNA edgeR QL按各范围/过滤臂全受检基因BH。','technical_delivery_20261002_ikzf2_stat1_v1/\nrevision_v2/final_release/data/\nS2_TF_associations.tsv；existing_RNA_DE_focused.tsv'],
 ['S5','受者供者伪bulk Base/Source DE；名义/FDR、Up/Down靶集分开，|logFC|≥0.5；受检且prior可用背景。原不足标准不改阈值。','human_first_20261004_v1/results/\nstage2_*；stage3_*；\nmouse_sender_20261002_v1/results/'],
 ['S6','独立候选发现家族和修订审计完整保留，不为正文三候选重算缩小BH。','phenotype_candidate_discovery_20261005_v1/\nresults/；deliverables/']],[1.1,7.6,8.6])
h('本次交付与溯源')
p('新目录：P/11_reports/phase_delivery_IKZF2_STAT1_20261005_v1。deliverables/ 保存 PDF、Word；results/ 保存两位TF的既有完整模型追溯表（含未支持与未计算状态）；provenance/ 保存原文件 SHA256 清单和交付回执；code/ 保存本次整理及排版代码。')
p('报告中的代表性表格用于阅读，不替代原完整模型。既有 human_first、technical_delivery、gap、independent_candidates 与 discovery 目录保留全部模型、ligand-target/背景成员、有限规格和原复现包。本次服务器补充ZIP只封装收束报告、来源表和整理代码。')
p('缩写：CWP=原病例标签；Control=对照；Core=排除CD4state的CD4范围；Primary=包含CD4state；AUC=调控子活性评分；ρ=偏秩相关；BH=Benjamini–Hochberg；FDR=假发现率；AM=肺泡巨噬；MLM=单核样巨噬。')
paragraphs=list(doc.paragraphs)+list(s.header.paragraphs)+list(s.footer.paragraphs)
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:paragraphs.extend(cell.paragraphs)
for pp in paragraphs:
 for rr in pp.runs:
  rr.font.name='Microsoft YaHei';rf=rr._element.get_or_add_rPr().rFonts
  for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Microsoft YaHei')
  for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
   if qn('w:'+a) in rf.attrib:del rf.attrib[qn('w:'+a)]
out=D/'CD4_IKZF2_STAT1_stage_delivery_zh.docx';doc.save(out)
print(out, out.stat().st_size)
