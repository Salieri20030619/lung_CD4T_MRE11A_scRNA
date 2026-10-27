from pathlib import Path
import json, pandas as pd, re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).parent;D=B/'manuscript_delivery'
E=json.loads((B/'revision_delivery/report_evidence_aggregate.json').read_text(encoding='utf8'))
F=json.loads((D/'added_evidence_aggregate.json').read_text(encoding='utf8'))
de=lambda k:pd.DataFrame(E[k]);df=lambda k:pd.DataFrame(F[k])
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.85);sec.right_margin=Cm(1.85)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(10.5);st.font.color.rgb=RGBColor.from_string('111111');st.paragraph_format.line_spacing=1.18;st.paragraph_format.space_after=Pt(5)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
for name,size in [('Title',18),('Subtitle',10),('Heading 1',13),('Heading 2',11.5),('Heading 3',10.5),('Caption',8.6)]:
 st=doc.styles[name];st.font.size=Pt(size)
 if name!='Caption':st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimHei')
doc.styles['Caption'].font.bold=False;doc.styles['Caption'].font.italic=False;doc.styles['Subtitle'].font.italic=False
for st in doc.styles:
 for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for n in ['Heading 1','Heading 2','Heading 3']:
 doc.styles[n].paragraph_format.space_before=Pt(11);doc.styles[n].paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.first_line_indent=Pt(21)
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
ftr=sec.footer.paragraphs[0];ftr.alignment=1
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');ftr._p.append(fld)
tables=[];figures=[];evidence_checks=[]
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t,level=2):doc.add_heading(t,level)
def n(x,d=4):
 if x is None or pd.isna(x):return '不可估计'
 if isinstance(x,str):return x
 if x!=0 and abs(x)<.001:return f'{x:.4g}'
 return f'{x:.{d}f}'
def table(title,heads,rows,widths=None,note=None,supp=False):
 idx=('S' if supp else '')+str(sum(t['supp']==supp for t in tables)+1);tables.append({'id':idx,'title':title,'supp':supp,'rows':len(rows)})
 pp=p(f'表{idx}　{title}','Caption');pp.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(heads));t.autofit=False;widths=widths or [17.2/len(heads)]*len(heads)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,v in zip(t.rows[0].cells,heads):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);pr=c._tc.get_or_add_tcPr()
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF4');pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.space_before=Pt(1.7);pp.paragraph_format.space_after=Pt(1.7);pp.paragraph_format.line_spacing=1.02
    if ri==0:pp.paragraph_format.keep_with_next=True
    for r in pp.runs:r.font.size=Pt(8.3);r.bold=ri==0
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 if note:p(note,'Caption')
 return idx
def fig(file,title,caption,supp=False):
 idx=('S' if supp else '')+str(sum(x['supp']==supp for x in figures)+1);figures.append({'id':idx,'title':title,'file':file,'supp':supp})
 pp=p('');pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.keep_with_next=True;pp.add_run().add_picture(str(D/'figures'/file),width=Cm(17.15))
 p(f'图{idx}　{title}。{caption}','Caption')
def textblock(s):
 for line in s.strip().split('\n'):
  if line.strip():p(line.strip())

title=p('尘肺患者 CD4 T 细胞转录状态、\n巨噬细胞通讯与 IKZF2、STAT1 调控证据','Title');title.alignment=1
pp=p('人类患者数据与小鼠硅尘模型的整合分析','Subtitle');pp.alignment=1
pp=p('数据及文献核对截至 2026 年 10 月 6 日','Caption');pp.alignment=1
h('摘要',1)
textblock('''目的：以人类患者为主要证据，明确疾病相关 CD4 T 细胞组成和内部转录程序，评估巨噬细胞信号与接收者疾病程序的吻合程度，并独立评价 IKZF2 与 STAT1 的调控证据。小鼠数据用于方向和亚群层面的对照，不替代人类疾病验证。
方法：复用经核对的单细胞对象、供者清单和既有分析。人类 RNA 队列含 19 名供者（10 名对照、9 名煤工尘肺患者），来源分布不平衡；不同分析按预定细胞数门槛保留相应供者。以供者为统计单位，分析 CD4 亚群比例、伪批量差异表达、GSEA、GSVA、UCell 及表达匹配程序评分。SCENIC 网络采用三个随机种子，并比较原始、剔除候选基因及剔除程序重叠基因后的活性。通讯分析区分配体–受体表达支持、NicheNet 配体活性和同供者关联；重采样保留完整供者行与来源–疾病分层。小鼠主队列为每组 3 个独立文库。
结果：人类 CD4 组成的基础模型可见名义差异，但来源校正改变部分结论。GZMA 亚群的调节性 T 细胞相关程序及 CCR7 亚群的抑制相关程序存在名义疾病关联，扩展程序家族的 BH 校正未通过。人类氧化磷酸化 GSEA 富集稳定：核心 CD4 的 KEGG 标准化富集分数为 2.774，FDR=2.06×10⁻¹³；剔除核糖体及线粒体编码基因后仍保留。相同供者的 GSVA 未有通路通过 BH 校正，整体线粒体应激及焦亡 UCell 评分也未见疾病差异。核心 CD4 中，去重叠基因后的 IKZF2 活性与焦亡执行相关评分在三个种子均负相关（n=17，ρ=−0.830、−0.838、−0.664；P=0.000446、0.000349、0.013365；原 FDR=0.02567、0.02883、0.11077）；背景筛选后 n=15，三个种子均未通过原 BH 校正。凋亡执行相关评分同向负相关，但三个种子的原 FDR 均未通过。人类 GZMA 亚群 IKZF2 RNA 的来源校正效应为 +0.04453（P=0.002109，FDR=0.025304），加入深度后 FDR=0.051376；STAT1 的人类 RNA 疾病差异未获支持。小鼠核心 CD4 的两候选 RNA 和调控子活性有上调证据，但并非所有亚群或独立时间点一致。人类核心和 IL7R 亚群的名义上调靶集中 IL15 排名靠前，FDR 靶集不足；GZMA 的 FDR 上调靶集可计算且出现不同候选。TNF、IL18 的配体关联、疾病程序和重采样区间未形成稳定的因果证据链。
结论：人类数据支持部分 CD4 转录状态、呼吸链程序及特定调控子–程序关联，但不能据此确认线粒体损伤、凋亡、焦亡或配体驱动的机制。IKZF2 与 STAT1 具有不同的证据组合，应作为两个独立候选保留。来源混杂、有效供者不足、靶集门槛和模型敏感性仍限制机制解释。''')
p('关键词：煤工尘肺；CD4 T 细胞；巨噬细胞；IKZF2；STAT1；GSEA；SCENIC；NicheNet','Caption')
h('1　研究背景与方法',1)
h('1.1　研究问题及证据层级')
textblock('''尘肺相关肺纤维化伴随持续的髓系炎症和适应性免疫变化。晶体硅尘模型中，NLRP3 炎症小体的遗传学证据支持其参与肺部炎症及纤维化，但该结论不能直接延伸为患者 CD4 T 细胞中的 IL18–转录因子机制。[1] 巨噬细胞外泌体影响 T 细胞平衡的研究，为跨细胞信号提出了独立的生物学依据；其 miRNA、体外共培养及成纤维细胞读数与本研究的单细胞配体预测处于不同证据层级。[5]
IKZF2 编码 Helios，其表达既与调节性 T 细胞有关，也可见于活化的常规 T 细胞，因此不将 IKZF2 高表达或相关程序增加等同于完整的抑制功能。[6] STAT1 与干扰素应答和辅助性 T 细胞分化有关；已有硅尘研究显示干扰素信号的作用具有细胞类型和组织环境依赖性。[2,8] 本研究分别考察疾病差异、调控子活性、程序关联和跨物种方向，不以一项阳性读数替代其余层面的验证。
分析沿“疾病表型—细胞内程序—调控证据—发送端与接收端吻合”的顺序组织。细胞组成、供者内表达、调控子 AUC、通路富集、配体活性和相关系数各自具有不同分母与统计零假设。三个计算种子属于算法敏感性，不是三个生物学重复；多个背景定义也不是独立队列。''')
h('1.2　数据来源、疾病标签与供者单位')
textblock('''人类最终 RNA 分析名单为 19 名供者。GSE135893 提供 2 名对照，GSE227136 提供 8 名对照和 2 名病例，GSE288887 提供 7 名病例。病例原始标签为 CWP（煤工尘肺），本报告不将其统一改写为严格诊断的矽肺，也不推断未核实的病程或严重度。病例与来源高度相关，来源内可比较的病例主要来自 GSE227136 的 2 人，因此来源校正不能消除所有不可辨识性。
人类上游肺细胞图谱覆盖 20 名供者；最终 CD4 分析的 19 人名单与该图谱不同，不能相互借用分母。全体 CD4 共 6523 个细胞，保留 CD4state；核心 CD4 共 5705 个细胞，排除该状态的 818 个细胞。亚群细胞数门槛、调控子可评分名单和背景筛选分别应用，未将固定的 17 人 AUC 名单强加给所有 RNA 或比例分析。
小鼠主要疾病对照来自 GSE314943，每组 3 个独立文库（Saline 与 Silica），按核心 CD4、Naive、ActCD4、Treg 分开。GSE183682 的 7 天与 56 天每条件仅 1 个文库，仅用于冻结调控子的方向对照，不计算重复不足条件下的疾病推断 P 值。其他已存队列的程序评分保留为聚合方向证据，未据此增加新候选。''')
table('分析人群与可解释范围',['数据/范围','独立单位及细胞覆盖','用于解释的限制'],[
['人类肺细胞图谱','20名供者；192575个细胞','只描述总体细胞构成；不等同于19人CD4名单'],
['人类全体/核心CD4','19名RNA供者；6523/5705细胞','全体含CD4state；核心排除CD4state'],
['人类伪批量通路','≥20细胞/供者后17人：8对照、9病例','来源校正；病例和来源强混杂'],
['人类调控子活性','全细胞分析17人；背景筛选后15人','细胞筛除同时改变供者覆盖'],
['小鼠主队列','3对照+3硅尘文库','Treg含低细胞文库，结果探索性'],
['小鼠独立时间点','7天/56天每条件1文库','仅描述方向，不作生物学重复推断']],[4.0,6.2,7.0])
h('1.3　注释复用、分母和发送端定义')
textblock('''复用已核验的 CD4-CCR7、CD4-GZMA、CD4-IL7R、Treg 和 CD4state 注释，不重新进行全量聚类。比例分母分别为同供者全体 CD4 或核心 CD4 总数；内部程序评分仅在相应亚群中汇总。比例为零的供者在有明确分母时仍保留，而程序评分需要实际接收者细胞；“有细胞”和“有至少20个细胞”的名单分别记录。
巨噬发送端包括肺泡巨噬细胞富集集合 MacAM、单核/巨噬混合集合 MonoMac，以及两者的并集 MacWide。MacAM 对应既有簇 0、2、11、13；MonoMac 对应簇 3、7，其 C1Q、MARCO、SPP1 与 FCN1 共同表达支持混合特征，不能整体视为纯肺泡巨噬细胞。MacWide 不构成独立重复。T/NK 来源仅用于竞争性来源的敏感性判断，不代替巨噬发送端。''')
h('1.4　供者模型、深度处理与多重检验')
textblock('''以供者或独立文库为推断单位，不将单细胞当作独立重复。[11] 比例及程序的基础模型为 Y~疾病；来源校正模型增加研究来源；来源与深度校正模型再增加相应细胞范围的 log1p(平均原始 UMI)。GSE227136 来源内模型仅比较同一来源的疾病组。固定名单上增加协变量与筛去背景较高细胞属于不同操作，分别报告。
本报告沿用各分析原有的深度定义，而非以同一个“深度校正”名称掩盖差异。新增供者均值 RNA/调控子疾病模型使用 log10(平均深度)；较早的调控子组间模型使用 log1p(平均深度)；调控子–程序相关采用秩变换后的深度残差化；UCell 质量校正采用标准化 log10(中位 UMI+1)和标准化线粒体比例。伪批量 edgeR 则以文库大小及 TMM 为偏置项，不额外加入供者平均深度。完整公式见统计补充。
报告效应、原始 P 和实际计算的 BH 结果。12 程序疾病模型的 BH 在细胞范围、背景筛选、模型和指标家族内跨程序计算；较窄程序面板的跨范围 BH 另行标识。调控子–程序关联保留原分析的“受检TF×程序及候选表达特征”比较家族，不把本报告仅展示的 IKZF2/STAT1 重新定义为校正家族。旧分析中未计算的组间 FDR 或小样本相关 P 明确写为未计算，不当作阴性；较新的可估计供者模型不套用“每组不足10人不计算”的旧限制。''')
h('1.5　伪批量差异表达与通路分析')
textblock('''按供者与细胞范围汇总原始计数。人类 GSEA 的既有排序来自来源校正 edgeR 准似然模型，使用 filterByExpr、过滤后重置文库大小、TMM 和稳健离散度估计。全体及核心 CD4 的排序分别包含 10348 和 10961 个受检基因，排序值为 sign(logFC)×√QLF。该历史通路模型与通讯差异分析保留完整原始文库偏置的实现不同，二者不互换。
人类基因集取自 msigdbr 的 GO Biological Process 与 KEGG Legacy，GSEA 使用 clusterProfiler 的 fgsea 实现，基因集大小10–500，保留全部 P 值，eps=0，随机种子20260813。[13] 既有日志确认 msigdbr 26.1.0、clusterProfiler 4.18.4；基因集内容保存在分析对象中，未获得可独立核实的 MSigDB 发布号，因此不据软件号推断数据库发布版本。BH 在细胞范围及数据库内完成。GSEA 是受检基因排序的富集统计，并非细胞死亡率或代谢通量测定。
呼吸链敏感性在相同疾病排序上依次剔除胞质核糖体基因 RPL/RPS（保留 RPS6K），再剔除线粒体核糖体 MRPL/MRPS 及 MT- 编码基因；不重新拟合差异模型。GSVA 使用同一批供者的 logCPM（prior.count=2）、Gaussian 核和10–500基因集，随后以来源加疾病的 limma 模型及稳健、趋势经验贝叶斯检验比较。[14] GSVA 表中的 logFC 字段实际代表评分差值，不作为 RNA 倍数变化。
小鼠 GSEA 以独立文库的 edgeR 排序为输入。Ensembl–Entrez 映射使用 org.Mm.eg.db 3.22.0：一对多 Entrez 映射剔除，多 Ensembl 对一 Entrez 时按六文库平均 logCPM 保留代表，避免按疾病差异选基因。GO 使用 GOALL 祖先映射及 GO.db 3.22.0；KEGG 使用已保存的小鼠 REST 数据。fgsea 1.36.2、R 4.5.2，基因集15–500，eps=0、种子42。完整基因、去背景/血红蛋白基因，以及进一步去免疫球蛋白排序分别报告。''')
h('1.6　线粒体、凋亡和焦亡相关评分')
textblock('''采用两套既有评分而不混称。表达匹配评分由签名基因的平均 log1p(CP10k) 减去匹配对照基因均值得到；按供者等权基线表达分为25个分位箱，每个签名基因最多选择20个表达接近的对照基因。控制池排除候选、全部签名、背景、线粒体编码及核糖体基因。至少3个可测且表达的签名基因、至少60%成员覆盖及20个控制基因才形成有效评分；跨队列绝对值不直接比较。
12个固定程序包括核编码氧化磷酸化、线粒体整合应激、氧化还原反应、线粒体自噬、炎症小体、焦亡执行相关、凋亡执行相关、T细胞活化、调节性、抑制性、初始状态及增殖。凋亡执行签名为 BAX、BAK1、APAF1、CASP9、CASP3、CASP7、CYCS；焦亡执行签名为 PYCARD、CASP1、GSDMD、IL1B、IL18。二者具有不同分子含义，分别计算与报告。IL1B/IL18 转录并不等同成熟分泌蛋白，CASP/GSDM 转录也不证明裂解或膜破裂。
人类 UCell 在既有单细胞对象上按秩评分，再汇总供者均值；[15]主分析每范围每供者≥20细胞，亚群审计≥10细胞。其9个面板含线粒体应激、氧化还原、自噬、炎症小体感知、经典及非经典焦亡、Gasdermin替代通路等，成员均可在对象中找到。UCell 的线粒体氧化还原签名含 FTL，共13基因；表达匹配版本不含 FTL，为12基因。UCell 面板没有独立凋亡执行签名，不能把含 CASP3 的替代 Gasdermin 面板改名为凋亡。历史脚本调用 UCell 默认接口，未显式冻结 maxRank 参数，不能把小鼠的显式参数移植给人类。
小鼠 UCell 使用原始计数、maxRank=1500、UCell 2.14.0、种子42，七个 GO 模块加手工焦亡执行面板（Casp1、Casp4、Gsdmd、Gsdme），去除历史研究目的基因 Mre11/Mre11a 后评分。每文库平均分的组差使用全部20种3+3标签分配的双侧精确置换，主要三个细胞范围的模块共同校正，低细胞 Treg 另设探索家族。既有 MRE11 相关分析独立保存，仅作历史对照，不改写为 IKZF2 或 STAT1 证据。''')
h('1.7　调控网络、去重叠与独立时间点评分')
textblock('''SCENIC 采用 pySCENIC 0.12.1、GRNBoost2、三个随机种子17/29/43，分别在人类和小鼠核心 CD4 的完整细胞及背景筛选后细胞中建立网络。[9] 每供者最多抽取250个细胞训练、不重复抽样扩增，随后对全部符合条件细胞评分。输入为完整原始文库归一化的 log1p(CP10k)，不用整合表达或插补值；基因需在至少10个或1%细胞中出现、与两个基序库相交且训练方差>10⁻⁸，不先按差异表达选择。
基序资源为 hg38/mm10 Refseq r80 mc_v10_clust 的转录起始位点±10 kb及上游500 bp/下游100 bp数据库，配套 v10 nr_clust 注释。富集阈值 NES=3、rank threshold=5000、AUC前5%、最小基因集20，AUCell不加权。网络边为共表达及基序支持的推断，图中权重不表示激活/抑制方向或实验验证。
去候选基因版本移除 IL1R2、PENK、TTC39C、FOXP1、TNFRSF9、LAG3、MRE11 七个既有候选与TF自身；去程序重叠版本进一步移除12程序的基因并集，保持相同表达排序与种子重新评分。该操作减少签名共享导致的数学关联，但不能排除所有共同上游因素。每个种子的原始与保留靶数均记录，不用最有利种子代替其余结果。调控子–程序关联在秩变换后残差化疾病、来源和接收者深度；完整、去候选和去重叠版本以及背景筛选分别呈现。
独立时间点分析冻结主小鼠队列的调控子成员，将其映射到 GSE183682 的可测基因后评分；不在该数据中重新推断网络。其每条件每时间点只有一个文库，任何同向变化均是描述性观察，不提供推断 P 值。''')
h('1.8　供者重采样与区间解释')
textblock('''对两个已冻结的探索关系分别计算：肺泡巨噬 IL18 与 IL7R 亚群 IKZF2，以及肺泡巨噬 TNF 与核心 CD4 STAT1。终端效应程序为去共享基因后的 PRF1、IFNG、NKG7。配体–TF、TF–程序、配体–程序是三项相关，不预设为因果链。
主要重采样按“来源×疾病”分层、整行有放回抽取供者，每个方案5000次；固定17人名单内的RNA与三个AUC种子共享抽样索引，以便配对比较。19人与17人的方案分别进行，单供者分层仍只能重复该供者。基础关联模型控制疾病、来源及两端秩深度；TNF关系另加入Treg比例作为敏感性。报告有效重采样的百分位95%区间，线性分位数插值，不用重复细胞扩大样本量。[12]
秩不足、残差自由度<3、恒定或接近零残差使一次重采样不可估计，不另抽样补足。有效次数少于1000不报告区间。未分层重采样用于说明来源分层对可估计性的影响，其区间均以有效抽样为条件；区间跨零不能写成稳定支持。三个相关方向同时符合的比例仅是重采样中的代数相容率，不是中介效应或因果概率。''')
h('1.9　巨噬细胞至 CD4 的通讯与 NicheNet')
textblock('''在人类所有经核验的接收亚群和巨噬来源上建立通讯证据。接收者疾病差异以供者伪批量计数拟合基础模型或来源校正模型，保留完整原始文库偏置和TMM；按原定阈值分别形成 P<0.05 或 FDR<0.05、|log₂FC|≥0.5 的上调及下调靶集。两端每供者≥20细胞；主表达门槛为至少两个病例供者中≥10%细胞检出，5%仅为表达敏感性。低检出不被解释为受体在生物学上不存在。
接收者背景为受检基因、至少两个合格供者中达到表达门槛的基因及配体–靶先验共同交集；背景≥500、靶集≥10、非靶背景≥10且候选配体≥2才计算配体活性。名义靶集不能替代未满足门槛的FDR靶集而沿用同一证据标签。配体–受体表达与复合物兼容性、配体活性、预测靶与真实疾病差异的重叠，以及同供者关联四层分别保存。[10]
人类使用 nichenetr 2.2.1.1 的人类经典先验（Zenodo 3260758）；小鼠使用v2小鼠先验（Zenodo 7074291）。记录物种、别名和重复基因映射，跨物种比较方向及同源基因，不比较两个先验下的绝对排名。校正AUPR用于排序，不是P值；排名第一但校正AUPR为负仍不构成正向解释能力。先验为无符号网络，下调靶集活性不能直接解释为配体抑制作用。
同供者关联另要求两组各至少两人、两端有效细胞数及残差自由度≥3，控制疾病、来源及两端秩深度；在接收范围、背景筛选、差异模型、靶集门槛和方向内，跨配体–发送端组合校正。MacWide作为来源并集报告，不能与其组成集合累计为重复证据。''')
