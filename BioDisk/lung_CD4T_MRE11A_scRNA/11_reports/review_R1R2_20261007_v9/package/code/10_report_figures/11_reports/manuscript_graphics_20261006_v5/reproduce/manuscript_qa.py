from pathlib import Path
import pymupdf as fitz,json,re,hashlib
from docx import Document
from PIL import Image,ImageOps,ImageDraw
D=Path('manuscript_delivery');Q=D/'QA';Q.mkdir(exist_ok=True)
pdf=fitz.open(D/'CD4_IKZF2_STAT1_integrated_manuscript_zh.pdf');doc=Document(D/'CD4_IKZF2_STAT1_integrated_manuscript_zh.docx')
texts=[];bounds=[];pages=[]
for i,page in enumerate(pdf):
 text=page.get_text();texts.append(text);bad=[]
 for b in page.get_text('dict')['blocks']:
  if 'lines' not in b:continue
  for line in b['lines']:
   for s in line['spans']:
    x0,y0,x1,y1=s['bbox']
    if x0<20 or x1>page.rect.width-20 or y0<15 or y1>page.rect.height-15:bad.append({'text':s['text'],'bbox':s['bbox']})
 if bad:bounds.append({'page':i+1,'items':bad})
 p=Q/f'page_{i+1:02}.png';page.get_pixmap(matrix=fitz.Matrix(1.15,1.15),alpha=False).save(p);pages.append(p)
full='\n'.join(texts);(Q/'extracted_text.txt').write_text(full,encoding='utf8')
docfull='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]+[p.text for s in doc.sections for p in s.header.paragraphs]+[p.text for s in doc.sections for p in s.footer.paragraphs])
blacklist=[r'(?i)scenic[ _-]?[1-4]',r'阶段[一二三四五12345]',r'(?i)step[ _-]?\d',r'(?i)front_half',r'(?i)SourceDep',r'(?i)DropOverlap',r'(?i)DropTarget',r'(?i)WithinG227',r'(?i)technical_delivery',r'(?i)mechanism_step',r'(?i)NoRP\+MT',r'(?i)SourceAdj',r'(?i)QCAdj',r'\bBase\b',r'\bBG1\b',r'\bSCENIC\d+\b',r'肾癌',r'肾透明细胞',r'Phase327',r'Bootstrap阶段',r'(?i)/mnt/']
hits={rx:re.findall(rx,docfull+'\n'+full) for rx in blacklist if re.search(rx,docfull+'\n'+full)}
empty=[]
for ti,t in enumerate(doc.tables):
 if any('FDR' in c.text or 'BH' in c.text for c in t.rows[0].cells):
  for j,c in enumerate(t.rows[0].cells):
   if c.text.strip() in ['FDR','BH'] and all(row.cells[j].text.strip() in ['','—','未计算','None','nan'] for row in t.rows[1:]):empty.append(ti+1)
for j in range(0,len(pages),6):
 im=Image.new('RGB',(1260,1780),'#D0D0D0');dr=ImageDraw.Draw(im)
 for k,path in enumerate(pages[j:j+6]):
  pg=ImageOps.contain(Image.open(path),(615,560));x=(k%2)*630+(615-pg.width)//2;y=(k//2)*590+22;im.paste(pg,(x,y));dr.text(((k%2)*630+10,(k//2)*590+5),f'Page {j+k+1}',fill='black')
 im.save(Q/f'contact_{j//6+1}.jpg',quality=88)
result={'pages':len(pdf),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'text_characters':len(full),'out_of_bounds':bounds,'internal_term_hits':hits,'empty_BH_columns':empty,'replacement_chars':full.count('\ufffd'),'page_text_lengths':[len(t) for t in texts],'sha256':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in D.glob('CD4_IKZF2_STAT1_integrated_manuscript_zh.*')}}
(Q/'QA_result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(result,ensure_ascii=True))
