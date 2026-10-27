from pathlib import Path
import fitz,json,hashlib,re
from PIL import Image,ImageDraw
from docx import Document
B=Path(__file__).parent;D=B/'revision_delivery';Q=D/'QA';Q.mkdir(exist_ok=True)
f=D/'CD4_IKZF2_STAT1_stage_delivery_zh.pdf';pdf=fitz.open(f);report=[];thumbs=[]
for i,pg in enumerate(pdf):
 text=pg.get_text();outside=[];fonts=set();bottom=0
 for b in pg.get_text('dict')['blocks']:
  for line in b.get('lines',[]):
   for sp in line['spans']:
    fonts.add(sp['font']);x0,y0,x1,y1=sp['bbox']
    if x0<0 or y0<0 or x1>pg.rect.width+.5 or y1>pg.rect.height+.5:outside.append(sp['text'])
    if y0>35 and y1<795:bottom=max(bottom,y1)
 pg.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(Q/f'page_{i+1:02d}.png')
 im=Image.open(Q/f'page_{i+1:02d}.png').convert('RGB');im.thumbnail((500,708));can=Image.new('RGB',(520,738),'#dce4e8');can.paste(im,((520-im.width)//2,8));ImageDraw.Draw(can).text((10,722),f'Page {i+1}',fill='black');thumbs.append(can)
 report.append(dict(page=i+1,characters=len(text),out_of_page=outside,fonts=sorted(fonts),body_bottom=round(bottom,1),beginning=text[:90]))
for offset in range(0,len(thumbs),6):
 sheet=Image.new('RGB',(1560,1476),'white')
 for j,im in enumerate(thumbs[offset:offset+6]):sheet.paste(im,((j%3)*520,(j//3)*738))
 sheet.save(Q/f'contact_{offset+1:02d}_{min(offset+6,len(thumbs)):02d}.jpg')
document=Document(D/'CD4_IKZF2_STAT1_stage_delivery_zh.docx');empty_fdr=[];empty_tables=[]
for i,t in enumerate(document.tables,1):
 if len(t.rows)<2:empty_tables.append(i)
 for col,c in enumerate(t.rows[0].cells):
  if 'FDR' in c.text:
   vals=[r.cells[col].text.strip() for r in t.rows[1:]]
   if all(v in ['','—','NA','未计算'] for v in vals):empty_fdr.append(i)
whole='\n'.join(pg.get_text() for pg in pdf)
result=dict(pages=len(pdf),pages_detail=report,sha256=hashlib.sha256(f.read_bytes()).hexdigest(),automated_bounds_pass=not any(x['out_of_page'] for x in report),empty_fdr_columns=empty_fdr,empty_tables=empty_tables,tables=len(document.tables),figures=len(document.inline_shapes),all_reference_markers=all(f'[{i}]' in whole for i in range(1,13)),visual_review='pending')
(Q/'QA_render.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');(Q/'report_text.txt').write_text(whole,encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k!='pages_detail'},ensure_ascii=False));print([(r['page'],r['characters'],r['body_bottom']) for r in report])
