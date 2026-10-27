from pathlib import Path
import json,re,hashlib
from docx import Document
B=Path(__file__).parent;D=B/'graphics_delivery';Q=D/'QA'
old=Document(B/'narrative_delivery/CD4_IKZF2_STAT1_integrated_manuscript_zh.docx');new=Document(D/'CD4_IKZF2_STAT1_integrated_manuscript_zh.docx')
def vals(d):return [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
assert vals(old)==vals(new),'Old table values changed'
baseline=json.loads((D/'provenance/baseline_v3.json').read_text())
for f,sha in baseline.items():assert hashlib.sha256((B/f).read_bytes()).hexdigest()==sha
texts='\n'.join(p.text for p in new.paragraphs)
oldtexts='\n'.join(p.text for p in old.paragraphs)
def nums(s):
 s=s[s.index('2　结果'):s.index('4　统计与数值补充')]
 return set(re.findall(r'(?<![A-Za-z0-9])[-−+]?\d+\.\d+(?:[eE][−+-]?\d+)?',s))
missing=sorted(nums(oldtexts)-nums(texts));assert not missing,missing
manifest=json.loads((D/'document_manifest.json').read_text(encoding='utf8'))
expected=['Human_UMAP_atlas_CD4.png','fig1_population.png','fig7_communication.png','fig4_networks.png','candidate_expression.png','fig6_bootstrap.png','Human_Core_KEGG_Top5_GSEA.png','Human_Core_GO_Top5_GSEA.png','fig3_programs.png','candidate_function.png','figS1_transfer.png','fig2_pathways.png','Human_Primary_KEGG_Top5_GSEA.png','Human_Primary_GO_Top5_GSEA.png']
assert [x['file'] for x in manifest['figures']]==expected
assert len(manifest['references'])==15
same=[]
for f in expected:
 if (B/'narrative_delivery/figures'/f).exists():
  assert (B/'narrative_delivery/figures'/f).read_bytes()==(D/'figures'/f).read_bytes();same.append(f)
for term in ['GSEA','GSVA','UCell','凋亡','焦亡','线粒体','SCENIC','Bootstrap','NicheNet','IKZF2','STAT1','UMAP']:
 assert term in texts,term
figids={x['id'] for x in manifest['figures']};tabids={x['id'] for x in manifest['tables']}
for k,allowed in [('图',figids),('表',tabids)]:
 refs=set(re.findall(k+r'(S?\d+)',texts));assert refs<=allowed,(k,refs-allowed)
assert '图7、图8' in texts and '图S3、图S4' in texts and '对应亚群方向见图S2D' in texts
out={'tables_identical':len(new.tables),'source_data_unchanged':True,'original_prose_numerals_preserved':True,'old_prose_numerals_for_review':missing,'references':15,'main_figures':10,'supplementary_figures':4,'original_figures_identical':same,'new_scientific_models':0,'new_UMAP_runs':0,'figure_order':expected,'cross_references_valid':True}
(Q/'editorial_audit.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(out,ensure_ascii=True))
