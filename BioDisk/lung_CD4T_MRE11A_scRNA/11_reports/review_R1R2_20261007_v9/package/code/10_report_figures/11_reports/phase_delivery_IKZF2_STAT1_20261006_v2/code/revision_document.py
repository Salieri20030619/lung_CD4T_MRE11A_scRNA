from pathlib import Path
import json,pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).parent;D=B/'revision_delivery'
E=json.loads((D/'report_evidence_aggregate.json').read_text(encoding='utf-8'))
df=lambda k:pd.DataFrame(E[k])
doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.55);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(10.5);st.paragraph_format.line_spacing=1.13;st.paragraph_format.space_after=Pt(6)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for name,size,color in [('Title',25,'173C50'),('Heading 1',17,'173C50'),('Heading 2',12.5,'167C87'),('Heading 3',11,'173C50'),('Caption',8.2,'51616A'),('Subtitle',11,'52616B')]:
 doc.styles[name].font.size=Pt(size);doc.styles[name].font.color.rgb=RGBColor.from_string(color)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15);doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
hdr=sec.header.paragraphs[0];hdr.text='肺部疾病单细胞研究  |  完整阶段交付 · 修订版';hdr.style='Caption'
ftr=sec.footer.paragraphs[0];ftr.alignment=2;ftr.add_run('2026-10-06  ·  ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');ftr._p.append(field)
tables=[];figures=[]
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t,level=2):doc.add_heading(t,level)
def num(x,d=4):
 if x is None or pd.isna(x):return '—'
 if isinstance(x,str):return x
 if x!=0 and abs(x)<.0001:return f'{x:.3g}'
 return f'{x:.{d}f}'
def signed(x,d=4):return '—' if x is None or pd.isna(x) else f'{x:+.{d}f}'
def table(title,heads,rows,widths=None,note=None):
 idx=len(tables)+1;tables.append(title);pp=p(f'表{idx}  {title}','Caption');pp.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(heads));t.style='Light Shading Accent 1';t.autofit=False;widths=widths or [17.4/len(heads)]*len(heads)
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,v in enumerate(heads):t.cell(0,i).text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i])
   for pp in c.paragraphs:
    pp.paragraph_format.line_spacing=1.03;pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(3)
    for r in pp.runs:r.font.size=Pt(8.5)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 if len(t.rows)<=6 and sum(len(c.text) for row in t.rows for c in row.cells)<1300:
  for row in t.rows[:-1]:
   for c in row.cells:
    for pp in c.paragraphs:pp.paragraph_format.keep_with_next=True
 if note:p(note,'Caption')
 return idx
def fig(name,title,caption,width=17.3):
 idx=len(figures)+1;figures.append(title);pp=p('');pp.paragraph_format.keep_with_next=True;pp.add_run().add_picture(str(D/name),width=Cm(width));p(f'图{idx}  {title}。{caption}','Caption');return idx
h('人类 CD4 T 细胞变化、\nIKZF2 / STAT1 与巨噬细胞信号',0)
p('完整阶段交付报告 · 背景、方法、SCENIC、Bootstrap 与 NicheNet 结果','Subtitle')
p('修订日期：2026年10月6日。复用已完成分析和冻结结果；未重跑网络、重抽样或候选筛选。')
h('核心结论与阅读路线')
p('本项目保留 IKZF2 和 STAT1 为两个独立候选。IKZF2 在人类特定 CD4 亚群具有来源校正 RNA 差异及调控子–程序关联；STAT1 在人 GZMA 亚群具有调控子活性–炎症小体相关程序关联，并有小鼠 RNA/活性证据作为对照。支持层级不同，不能写成二者在人类所有疾病模型均通过，也不要求互为上下游。')
table('可报告结论与对应条件',['问题','已有积极结果','同时保留的边界'],[
['患者 CD4 是否改变','组成基础模型及 GZMA、CCR7 内部程序有疾病线索。','来源调整改变部分结论；亚群、分母和名单分开。'],
['IKZF2','人GZMA RNA Source：n=12，P=0.002109、FDR=0.025304；IL7R程序关联三种子名义支持。','GZMA Base P=0.4442；IL7R关联为探索性，原未计算FDR。'],
['STAT1','人GZMA活性–Inflam程序三种子P=0.0086–0.0309；鼠Core/ActCD4 RNA名义支持。','人程序关联原FDR未通过；人疾病RNA差异不支持。'],
['巨噬通讯','人IL15对部分名义靶集排名较高；鼠Core IL15排名3；若干配对关联保留。','排名不是P；低检出、来源混杂和负活动分数均保留。'],
['Bootstrap','量化供者抽样下的区间、符号稳定性及调整差异。','条件于现有来源构成，不证明中介或级联。']],[3,7.3,7.1])
p('阅读路线：背景 → 数据流转 → 模型词典 → 患者表型与RNA → SCENIC1–4 → Bootstrap → 人鼠NicheNet → 补充候选及文献比较。附录提供模型家族、完整规格索引和核心敏感性。')
h('1  研究背景与科学问题',1)
h('1.1  矽尘先天炎症与 CD4 异质性')
p('矽尘相关肺损伤涉及先天炎症与适应性免疫。Cassel 等的鼠巨噬和遗传缺陷实验支持 NALP3/NLRP3 炎症小体在实验性矽肺中的作用，为 IL1β/IL18 等信号提供疾病背景；这并未建立 IL18 到本项目 CD4 亚群或两个 TF 的路径。[1] Ding 等报道巨噬细胞外泌体 miR125a-5p 与 T 细胞失衡及成纤维细胞变化，说明巨噬–T 相互作用已有邻近机制研究；它与本项目配体/TF关联属于不同证据体系。[5]')
p('CD4 不能视为单一功能群。You 等在鼠矽肺模型中区分 CD69⁺CD103⁻ 与 CD103⁺ 驻留记忆 CD4 细胞，观察到不同致病/免疫抑制特征及 IL7 相关维持机制。[3] 这支持按亚群研究，但不允许直接把本项目的人 CCR7、GZMA、IL7R 或 CD4state 改名为该论文细胞群，也不能把 IL7 结果替换为 IL15 机制。')
h('1.2  IKZF2 与 STAT1 的独立生物学依据')
p('IKZF2 编码 HELIOS，其关联可涉及 Treg 和常规 T 细胞状态；活化/增殖也可诱导 HELIOS，因此单个 IKZF2 信号不足以识别胸腺来源 Treg 或证明完整免疫抑制功能。[6] 新近人多组织研究在婴儿脾 T 细胞 CRISPR 实验中支持 HELIOS 对部分效应程序的限制，为“IKZF2–效应状态”提供邻近机制依据；关键扰动不是成人 CWP 肺 CD4，不能作为本项目直接验证。[7]')
p('STAT1 与 IFNγ、T-bet 和初始 CD4 分化反应有关，经典鼠实验提供了机制依据。[8] 但疾病方向依赖细胞、阶段和干预：Peng 等在人 BALF 及鼠模型中发现 IFNγ 信号受损，通过全身/髓系受体缺失和巨噬实验显示信号不足可加重病变；这不是 CD4 特异 STAT1 敲除。[2] DHI 药物研究观察到鼠早期全肺 pSTAT1/pSTAT3 和 Th 反应降低，属于多靶点干预，不能证明 STAT1 在所有细胞中均促纤维化。[4]')
p('本报告分别评价两位候选的患者表型/程序关联。巨噬信号用于免疫环境解释及后续优先级。旧 IL18–IKZF2、TNF–STAT1 三段模型保留为已开展探索，不作为两个候选必须通过的串联条件。')
h('1.3  本项目新增观察的定位')
p('本阶段关注：患者哪些已核验 CD4 亚群的组成和内部程序改变；IKZF2/STAT1 的 RNA、预测调控子活性与程序关联分别支持什么；巨噬表达、配体–靶预测和同供者关联是否与这些程序相呼应。新增价值是本队列、特定亚群和明确模型下的分层关联及候选优先级，不宣称首次发现矽肺巨噬–T 通讯或已证明保护/致病机制。')
h('2  数据、注释与供者流转',1)
p('疾病统计以供者或独立动物文库为单位。每个供者的细胞用于估计其表达、比例和程序均值，不当作独立重复。该原则与单细胞差异分析中重视生物学重复一致，但不能消除来源混杂。[11]')
table('队列与用途',['数据','独立单位 / 原标签','用途'],[
['人肺整合RNA','19供者：10 Control、9 CWP；G135/G227/G288。','主要表型、RNA与通讯。严重度、暴露强度等未全面核验。'],
['人SCENIC/AUC','All：17供者（8/9）；BG1：15供者。','沿用上游Core≥20冻结名单；缺失AUC不补零。'],
['鼠GSE314943','3 Saline / 3 Silica独立文库。','RNA、SCENIC、NicheNet；Core与ActCD4分开。'],
['鼠GSE183682','NS/Silica，7和56天，每时间×组仅1文库。','SCENIC4迁移与方向；不提供正式组间P。'],
['人GSE174725','BALF：2 Exposed / 3 Silicosis。','独立背景和既有固定验证。'],
['人GSE296682','两位患者，无对照。','发送端定位和受体覆盖，不能检验疾病差异。']],[3.4,6.9,7.1])
p('G135=GSE135893，G227=GSE227136，G288=GSE288887。主要患者沿用整合对象的 CWP（煤工尘肺）标签，不改写为严格 silicosis。肺组织与 BALF 采样不同；患者、动物及体外文献分别解释。')
table('来源与疾病设计',['来源','Control','CWP','解释'],[['G135',2,0,'仅对照'],['G227',8,2,'唯一两组均有的来源；病例仅2'],['G288',0,7,'仅CWP']],[3,2,2,10.4])
p('Source 是模型中的来源变量，不是新病人分组。G288无对照、G135无病例限制疾病/来源的区分；WithinG227仅在混合来源内部比较，仍不能凭2病例替代独立大队列。')
fig('R01_flow.png','名单与范围流转','Ctrl=Control；BG1=固定背景UMI比例≤1%。亚群n指Cells20供者；Bootstrap的RNA19/匹配17人是另一个固定名单。箭头表示分析范围，非重新聚类。来源[A1,A2,A4,A5]。')
rows=[];cov=df('scope_coverage');cov=cov[(cov['filter']=='All')&(cov.lineage=='CD4')]
for scope in ['CD4-CCR7','CD4-GZMA','CD4-IL7R','Treg','CD4state']:
 q=cov[cov.scope==scope].set_index('group')
 rows.append([scope.replace('CD4-',''),int(q.cells.sum()),'/'.join(str(int(q.loc[g,'positive_donors'])) for g in ['Control','CWP']),'/'.join(str(int(q.loc[g,'donors20'])) for g in ['Control','CWP'])])
table('全部CD4亚群覆盖（All）',['范围','细胞数','有细胞供者：对照/CWP','Cells20：对照/CWP'],rows,[3.8,2.8,5.4,5.4],note='Primary：6523细胞；Core排除CD4state，RNA19中5705细胞。SCENIC All冻结后5684细胞。Primary/Core伪bulk均17人（8/9）。[A1,A2]')
p('两位对照仅3和18个Core细胞，在更早上游门槛被排除，形成AUC 19→17；BG1下另两病例Core细胞93/79降为10/2，形成17→15。IL7R固定背景分析从AUC17再因一位对照只有10个IL7R细胞而用16人。减少源于声明的门槛，不是按P值删供者。[A2,A3]')
h('3  模型词典：Base、Source、SourceDep不是病人分组',1)
h('3.1  读数与效应尺度')
table('表达、活性、程序和比例的区别',['读数','实际处理','方向与单位'],[
['edgeR RNA','同供者同范围counts相加，负二项QL模型，TMM及文库offset。','log2FC>0为病例计数表达升高；P/FDR属该模型。'],
['旧均值RNA OLS','单细胞CP10k先在供者内平均，再log1p。','系数是log1p(平均CP10k)差，不是log2FC。'],
['20261005 TF RNA OLS','每细胞先log1p(CP10k)，再在供者内平均。','系数是平均log1p CP10k差；与上一行不同。'],
['SCENIC / AUCell','预测regulon靶基因在细胞表达排名前端的富集，供者内平均。','AUC差不是TF RNA、蛋白活性、结合或KO效应。'],
['12冻结程序','成员平均log1p CP10k减匹配表达背景均值。','相对转录程序；TReg/PyroExec/Mitophagy不等于功能。'],
['front-half五面板','去两个TF三seed靶union后，基因log1p CP10k在冻结Core细胞标准化，再平均。','与匹配背景程序不同；Bootstrap复用State_Effector。'],
['组成比例','亚群细胞数除以声明的Primary/Core CD4总数。','系数为比例差，不是组织绝对细胞丰度。']],[3.4,7,7])
h('3.2  逐分析族公式')
p('符号：G=病例指示变量（人CWP/鼠Silica）；S=来源类别；Lᵣ/Lₛ=供者在接收/发送范围内每细胞原始总UMI的均值；r(x)=平均秩。βG>0表示病例结局升高；ρ>0表示两读数在该条件下同向，不是疾病效应。')
table('公式、变量和固定名单',['分析族 / 模型','实际设计','输出与名单'],[
['人表型OLS','Base：Y~1+G；Source加S；SourceDep再加log1p(Lᵣ)。比例用分母范围深度，程序用其亚群深度。','同scope/roster固定样本；βG、t检验P、95%CI。[A1]'],
['20261005 TF RNA/AUC','Base：Y~1+G；Source加S；SourceDep再加log10(Lᵣ)。','Cells20；每组≥2、设计可识别；βG。RNA为平均log1p CP10k。[A3]'],
['旧SCENIC2疾病AUC','鼠Group：Y~1+G；Depth加log1p(Lᵣ)。人Source加S；SourceDep再加log1p(Lᵣ)。Comp/SourceComp另加K−1亚群比例。','Core和亚群分别，Cells20；不是上一行的log10定义。[A4]'],
['旧SCENIC2 TF–程序','两端秩对D残差化再算Pearson；人SourceDep：D=[1,G,S,r(Lᵣ)]。','偏秩相关ρ；旧规则n<10不提供相关P。[A4]'],
['人NicheNet DE','Base：counts~G；Source：counts~G+S；原完整伪bulk文库×TMM作offset。','不额外加入平均深度。Base/Source是疾病靶集DE标签。[A6]'],
['人NicheNet配体–程序','D=[1,G,S,r(Lₛ),r(Lᵣ)]，两端秩分别残差化。','两端≥20细胞、每组≥2、df≥3。即使靶集标Base，相关也含两端深度和来源。[A6]'],
['Bootstrap B / B_Treg','B：D=[1,G,S,r(Lᵣ),r(Lₛ)]；B_Treg另加r(Treg_fraction)，仅TNF线。每次重算秩并标准化设计列。','按冻结19/17人整行配对抽；不重建AUC/程序。[A5]'],
['旧sender-TF','SourceDep：D=[1,G,S,r(Lᵣ)]；BothEpi另加发送深度及上皮负担（秩）。','额外发送变量仅用于发送关联，不是所有SourceDep都包含。[A4,A7]']],[4,8,5.4])
p('偏秩相关的自由度为n−rank(D)−1，疾病OLS为n−rank(X)。P按原模型保留：包括t近似及明确标注的精确置换。后续stepwise诊断已对小n可估计模型计算P，未沿用旧“每组不足10就拒绝全部计算”的规则；旧未计算项则保持原身份。')
h('3.3  过滤、名单及调控子版本')
table('关键标签释义',['标签','含义'],[
['Primary / Core','人Primary含CD4state；Core排除它。鼠Core与ActCD4也分别报告。'],
['PositiveCells / Cells20','当前可用输入中该范围>0 / ≥20细胞的供者。PositiveCells不是TF阳性细胞，也不追回上游排除的AUC供者。'],
['All / BG1','原合格细胞 / 再按固定背景UMI比例≤1%筛CD4细胞；巨噬端不自动应用该过滤。'],
['旧Full / NoBG','原edgeR基因/归一化臂：背景和血红蛋白特征在表达过滤和TMM前移除；不是删细胞，不等同All/BG1。'],
['Original / DropTarget / DropOverlap','Original复用SCENIC1；DropTarget去旧7候选及TF自身；DropOverlap再去冻结模块成员union，同排名宇宙重算AUCell。'],
['NicheNet Full / Drop7','完整目标/背景 / 从两者同时去旧7候选；不等于SCENIC DropOverlap。'],
['seed17/29/43','训练抽样和排名随机性敏感性，不是独立生物重复。'],
['WithinG227','仅G227内Y~G，不再加来源项；全来源仅2病例。'],
['校正 / 筛除','Base→Source→SourceDep在固定名单上增加变量；BG1、Cells20及去主导供者是另一个筛选问题。']],[4.6,12.8])
h('3.4  P、FDR 与 Bootstrap 区间')
p('名义P<0.05的有限正式关联保留，并列效应、n、模型。edgeR的BH覆盖每范围/过滤/模型全部受检基因；20261005 TF疾病模型按readout/variant/seed/filter/model跨scope与TF；旧TF–程序相关按species/filter/seed/variant/scope/model跨TF–feature。正文少数候选不另缩小家族。')
p('旧M2因当时每组≥10规则未计算FDR，后续保留同一P并修复BH；本报告明确使用12程序/范围家族或人类优先四程序跨范围家族。整表未计算FDR则不列空列，仅脚注说明原因；已计算但未通过的FDR保留。Bootstrap只有原效应、CI、方向比例及有效次数，没有新P/FDR。')

h('4  患者 CD4 表型：组成与内部程序分开',1)
h('4.1  组成变化')
p('组成回答的是“CD4内部份额是否变化”，不是组织绝对丰度。Primary和Core各有自己的分母；基础模型沿用全部19供者。三个名义结果见下表，来源调整后均未保留名义支持，因此可作为病例组成线索，不称为来源独立的稳定特征。')
q=df('stage1_models');q=q[(q.kind=='fraction')&(q.denominator=='Primary')&(q.model=='Base')]
table('Primary分母下的全部亚群比例比较',['亚群','n（CWP/对照）','比例差βG','P','BH FDR'],[[r.scope,'19（9/10）',signed(r.effect),num(r.P),num(r.FDR)] for r in q.itertuples()],[3.6,4,3.1,3.3,3.4],note='模型：fraction~G，正值表示CWP份额较高；BH在同一分母/模型的5亚群内。Core分母的CCR7差−0.0995（P=0.0442、FDR=0.0885），Treg差+0.1327（P=0.0149、FDR=0.0596）另存，不能与Primary比例混用。[A1]')
h('4.2  亚群内部程序')
p('12个冻结转录面板覆盖T细胞状态与线粒体/细胞死亡相关探索；不是统一官方通路集合。评分为成员均值减固定匹配背景，成员与背景保持冻结。先展示全部亚群的4个基础免疫程序，避免预先只关注GZMA或CCR7。')
fig('R02_CD4_programs.png','全部CD4范围的基础免疫程序','All/Cells20、Source模型Y~G+S。色值为调整后CWP−Control分数差；*仅名义P<0.05；NA表示不可估计。Primary/Core n17，CCR7 n9、GZMA n12、IL7R n16、Treg n10；CD4state仅1位合格对照而不足。TAct=活化相关、TReg=调节相关、Inhib=抑制标记相关、Naive=初始样程序。不是蛋白或功能检验。[A1]')
q=df('stage1_models');z=q[(q.kind=='program')&(q.roster=='Cells20')]
spec=[('Subtype:CD4-GZMA','score_TReg','Source'),('Subtype:CD4-GZMA','score_TReg','SourceDep'),('Subtype:CD4-GZMA','score_TReg','WithinG227'),('Subtype:CD4-GZMA','score_Inhib','Source'),('Subtype:CD4-CCR7','score_Inhib','SourceDep')]
rows=[]
for scope,metric,model in spec:
 r=z[(z.scope==scope)&(z.metric==metric)&(z.model==model)].iloc[0]
 rows.append([scope.replace('Subtype:CD4-','')+'/'+metric.replace('score_',''),model,f'{r.n}（{r.n_case}/{r.n_control}）',signed(r.effect),num(r.P),num(r.legacy_FDR)])
table('明确亚群内保留的疾病程序线索',['亚群/程序','模型','n（病例/对照）','βG','P','12程序FDR'],rows,[3.5,2.6,3.5,2.3,2.7,2.8],note='FDR使用已修复的旧12程序/范围/过滤/模型家族，原P不变；不与四程序跨范围家族混用。GZMA TReg基础模型β=+0.0476、P=0.000425，四程序跨范围BH=0.01019；在扩展12程序跨范围家族为0.03057。这些是不同家族，不是P被改写。[A1]')
p('GZMA TReg相关转录程序的结果支持有限疾病线索，不等同GZMA群被确认为功能Treg。CCR7 Inhib的名义证据依赖深度调整。近期Core/IL7R五小面板30项疾病比较没有名义支持，不能把这些面板的配体相关直接称为疾病表型。')
h('5  IKZF2 与 STAT1：先说明RNA证据',1)
h('5.1  人类既有伪bulk与新亚群模型')
p('原人Primary/Core来源校正edgeR已被原counts精确复现，不能用后续不同均值OLS的结果覆盖。两种范围均17人（9CWP/8对照），Core为5684细胞。原IKZF2 RNA有名义差异；STAT1没有。')
q=df('human_legacy_edgeR_audit')
table('原人类来源校正edgeR（独立于下表OLS）',['范围/基因','n','log2FC','P','原全基因FDR'],[[r.scope+'/'+r.gene,17,signed(r.logFC_original),num(r.PValue_original),num(r.FDR_original)] for r in q.itertuples()],[4.8,1.6,3.3,3.7,4],note='设计~Series+Group；系数CWP−Control。伪bulk计数QL模型，不与平均RNA OLS拼接。原/精确重拟合P最大差约5.55×10⁻¹⁶。[A2]')
q=df('human_TF_disease');q=q[(q.readout=='RNA')&(q.scope=='CD4-GZMA')&(q.tf=='IKZF2')]
table('IKZF2：人GZMA的20261005 RNA OLS',['模型','n（病例/对照）','效应及95%CI','P','该家族FDR'],[[r.model,f'{r.n}（{r.n_case}/{r.n_control}）',f'{signed(r.effect)} [{signed(r.CI_low)}, {signed(r.CI_high)}]',num(r.P,6),num(r.FDR,6)] for r in q.itertuples()],[2.7,3.5,5.3,2.9,3],note='结局为供者平均log1p CP10k；Source加来源，SourceDep另加log10平均深度。BH在readout/variant/seed/filter/model内跨范围及2TF，所示家族m=12。[A3]')
p('GZMA Source通过对应FDR；SourceDep仍名义支持而FDR=0.051376；Base不支持。结论依赖来源调整且RNA稀疏（既有审计29/2554细胞检出IKZF2），因此正向发现明确写为“该模型条件下的人GZMA RNA关联”。人STAT1 RNA在All可估计的CCR7/GZMA/IL7R/Treg/Core/Primary、Base/Source/SourceDep中均未达名义，不把程序相关替代RNA差异。')
h('5.2  小鼠：Core与ActCD4都要呈现')
q=df('old_RNA_DE');q=q[(q.species=='Mouse')&(q['filter']=='All')]
table('鼠GSE314943既有RNA差异（3+3文库）',['范围/基因','log2FC','P','原全基因FDR'],[[r.scope+'/'+r.gene,signed(r.logFC),num(r.PValue,6),num(r.FDR,6)] for r in q.itertuples()],[5.4,4,4,4],note='mechanism_step3的edgeR QL组别模型，保留全库offset，All臂；BG1另列服务器。鼠Core中的两基因均为名义上调；ActCD4中Ikzf2不支持，Stat1名义支持。[A2,A6]')
p('Core包含细胞组成影响的可能性；ActCD4是更窄定位。不能用ActCD4的Ikzf2 RNA不支持否定整个Core，也不能用Core阳性证明ActCD4内机制。RNA与AUC读数不同，下一节保留其不同结果。')
h('6  SCENIC：从网络推断到独立程序关联',1)
h('6.1  为什么做、如何得到调控子活性')
p('SCENIC用于提出共表达与motif支持的调控子，再以AUCell检查这些靶基因在细胞表达排名中的富集。三个步骤依次是网络推断、motif筛选、活性评分；它提供转录调控假说，不是TF蛋白活性、真实结合或扰动结果。[9] 人和鼠分别构建，不把小鼠基因改大写就作为人网络。')
table('本项目SCENIC1实际参数',['步骤','已核验设置'],[
['输入','人冻结Step7 Core、小鼠GSE314943 Core；All与BG1独立网络。原全库CP10k、log1p，不用integrated assay、插补或按疾病差异预选基因。'],
['基因/训练细胞','基因检出≥max(10,1%保留细胞)，与两套motif数据库交集；训练子集中方差>1e−8。每供者最多250细胞，少者不补抽；所有保留细胞评分。'],
['软件/随机性','pySCENIC 0.12.1；GRNBoost2；3种子17/29/43，8 workers；共2物种×2过滤×3seed=12次运行。'],
['cisTarget资源','人hg38、鼠mm10，RefSeq r80 / mc_v10_clust；窗口10kb上游/10kb下游与500bp上游/100bp下游；v10nr motif注释及对应TF列表；资源校验留档。'],
['motif/评分参数','原CLI未另改默认项：motif NES阈值3.0、rank threshold5000、恢复曲线AUC前5%、建模块min_genes20；这些从当时环境CLI源码核对。'],
['解释','seed边支持≥2/3表示算法重复性，不是独立验证。网络箭头不代表已验证结合/因果或方向。']],[3.3,14.1])
q=df('SCENIC_runs');rows=[]
for (sp,arm),v in q.groupby(['species','filter'],sort=False):
 rows.append([sp,arm,int(v.donors.iloc[0]),int(v.cells.iloc[0]),'/'.join(str(int(v[v.seed==s].regulons.iloc[0])) for s in [17,29,43])])
table('SCENIC1完成规模',['物种','过滤','供者/文库','评分细胞','调控子数：s17/s29/s43'],rows,[2.6,2.2,3.6,3.6,5.4],note='这些是实际运行规模，不是显著发现数。人All为8对照/9CWP；鼠两臂均3+3。[A4]')
fig('F06_regulon_network_display.png','已有motif支持网络的固定展示子集','复用原图F06：All、seed17，每个TF按原边权重选12条DropOverlap保留边，展示规则与P无关。蓝=IKZF2，红=STAT1；上排鼠Core（6文库/2311细胞），下排人Core（17供者/5684细胞）。箭头为预测TF–靶关系，无因果符号；完整边及三seed/两过滤在服务器。靶名单可形成进一步假说，不能据网络图认定基因KO效果。[A4]。',16.5)
h('6.2  DropOverlap为何必要、保留了什么')
p('若TF调控子与结局程序直接共享基因，相关可能部分由同一信号重复计分造成。SCENIC2依次剔除旧7候选与TF自身（DropTarget），再剔除冻结12程序成员union（DropOverlap）。它只改变评分成员，不重新推断网络；在同一表达/排名宇宙以同seed重算未加权AUCell。非空集合均保留，空集合显式标记，不按显著性改阈值。')
q=df('S2_regulon_exclusion_coverage');q=q[(q.variant=='DropOverlap')&(q['filter']=='All')]
table('两候选All调控子的成员保留',['物种/TF','seed17 原→保留','seed29 原→保留','seed43 原→保留'],[[sp+'/'+tf]+[str(int(q[(q.species==sp)&(q.tf==tf)&(q.seed==s)].original_genes.iloc[0]))+'→'+str(int(q[(q.species==sp)&(q.tf==tf)&(q.seed==s)].retained_genes.iloc[0])) for s in [17,29,43]] for sp in ['Human','Mouse'] for tf in ['IKZF2','STAT1']],[4.2,4.4,4.4,4.4],note='BG1及DropTarget完整覆盖表另存；种子间成员数量变化是网络随机性的一部分，不能作为生物学样本变化。[A4]')
h('6.3  疾病AUC：鼠有名义线索，人Core未形成同样疾病差异')
q=df('old_TF_group');q=q[(q['filter']=='All')&(q.variant=='DropOverlap')&(q.species=='Mouse')&(q.model=='Group')&q.scope.isin(['Core','ActCD4'])]
rows=[]
for (sc,tf),v in q.groupby(['scope','tf']):
 rows.append([sc+'/'+tf,'6（3/3）']+[f'{signed(v[v.seed==s].coef.iloc[0])}; P={num(v[v.seed==s].p.iloc[0],6)}' for s in [17,29,43]])
table('鼠DropOverlap AUC疾病效应：全部三种子',['范围/TF','n（病例/对照）','seed17 β/P','seed29 β/P','seed43 β/P'],rows,[3.8,3,3.5,3.5,3.6],note='Group模型Y~G；AUC原尺度差。旧组间BH要求每组≥10，故本整表原未计算FDR，删除空列。各P是正式原模型名义P。[A4]')
p('鼠Core中两TF三个种子均名义升高。ActCD4的IKZF2 AUC也名义升高，与其RNA不显著并不矛盾；STAT1 ActCD4在seed29/43名义支持，seed17不支持。加入深度及TregNoTF身份评分的48个已有模型仍分开保存。All ActCD4 IKZF2在DepthTreg三seed P=0.003335/0.003683/0.014994；这类模型只有2个残差自由度，不能作为大样本确认。')
q=df('old_TF_group');q=q[(q.species=='Human')&(q['filter']=='All')&(q.scope=='Core')&(q.variant=='DropOverlap')&(q.model.isin(['Source','SourceDep']))]
rows=[]
for (tf,model),v in q.groupby(['tf','model']):
 rows.append([tf+'/'+model,'17（9/8）']+[f'{signed(v[v.seed==s].coef.iloc[0])}; P={num(v[v.seed==s].p.iloc[0])}' for s in [17,29,43]])
table('人Core DropOverlap AUC疾病效应：全部三种子',['TF/模型','n（病例/对照）','seed17 β/P','seed29 β/P','seed43 β/P'],rows,[4,3,3.5,3.5,3.4],note='旧SCENIC2模型：Source为G+S，SourceDep另加log1p平均原始文库大小；本整表原未计算FDR。不是20261005新模型。[A4]')
p('人Core两TF的Source/SourceDep均未达名义；STAT1的负向只是描述，不能称显著反向。不同物种/范围的疾病AUC不是统一“通过/不通过”开关；患者亚群程序关联仍可独立评价。')
h('6.4  两位候选的患者程序关联')
q=df('IL7R_IKZF2_excess');q=q[(q.endpoint=='excess')&(q.variant=='DropOverlap')]
table('IKZF2：IL7R固定背景扣除程序关联',['seed','n（CWP/对照）','ρ','P','近似95%CI'],[[r.seed,'16（9/7）',signed(r.rho),num(r.P,6),f'[{signed(r.CI_low)}, {signed(r.CI_high)}]'] for r in q.itertuples()],[2.1,4,3.2,3.4,4.7],note='调整G+S+r(Lᵣ)；excess=冻结疾病上调程序−3组固定匹配背景平均。原未计算FDR，故不列空列。Original三seed同向、P=0.0023–0.0292；不是独立患者复制。[A3]')
p('负相关支持IKZF2活性与该IL7R相对程序的独立关联，但与GZMA RNA属于不同问题。原/匹配背景分解已提示一般状态和背景对评分的贡献，不能直接称为保护效应。')
q=df('old_TF_program');q=q[(q.species=='Human')&(q.tf=='STAT1')&(q.scope=='CD4-GZMA')&(q.feature=='score_Inflam')]
table('STAT1：GZMA炎症小体相关程序关联',['seed','n（CWP/对照）','ρ','P','原家族FDR'],[[r.seed,'12（5/7）',signed(r.rho),num(r.p,6),num(r.FDR,6)] for r in q.itertuples()],[2,4,3.2,3.8,4.4],note='All、DropOverlap、旧SourceDep偏秩模型D=[1,G,S,r(Lᵣ)]。Inflam成员：NLRP3/AIM2/NLRC4/PYCARD/CASP1。原家族跨全部TF–feature对，未针对STAT1缩小。[A4]')
p('三seed均名义正相关而原FDR未通过。Core同模型仅seed17名义支持（ρ=0.7313、P=0.004503、FDR=0.060743），seed29/43 P=0.1550/0.1805。程序关联并不证明该程序已在人疾病中显著改变，也不证明炎症小体组装或焦亡发生。')
h('6.5  SCENIC3/4：身份、发送端与外部方向边界')
p('SCENIC3未重跑网络：鼠ActCD4用不含测试TF/候选的TregNoTF评分作身份调整；FOXP3未检出层只作描述，不能命名为Treg阴性身份。人髓系工作注释以既有标记审阅后确定MacAM/MonoMac，并审计上皮负担、主导供者和受体复合物。其目的分别是辨认组成影响与发送端可解释性，不是新增疾病亚型。')
p('SCENIC4将鼠冻结调控子迁移到GSE183682，统一发现与外部的测量/排名宇宙并重新评分，没有外部重新训练或按P选基因。12个TF/seed/过滤组合的DropOverlap靶基因在共同特征中覆盖率均100%；这只说明可评分，不证明外部疾病效应复制。')
fig('R05_SCENIC4.png','SCENIC4外部时间方向','All；色值为Silica−NS供者平均AUC差。Core细胞：7天271/393，56天416/333；ActCD4：7天60/95，56天135/108（顺序NS/Silica）。每时间×条件仅1文库，故无P/FDR。Core方向总体一致；ActCD4第7天IKZF2三seed为负、STAT1仅seed43为正，第56天均正。不能以时间点当独立重复。[A4]')
p('对主问题的贡献：SCENIC支持在RNA之外评估独立转录程序，并明确了范围、种子和时期依赖。它保留候选而不代替患者疾病比较，也不要求构成配体→TF级联。')

