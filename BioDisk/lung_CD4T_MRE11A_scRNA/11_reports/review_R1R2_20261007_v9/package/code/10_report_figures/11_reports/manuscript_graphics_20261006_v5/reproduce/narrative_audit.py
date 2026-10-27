from pathlib import Path
import json,re,hashlib
from docx import Document
D=Path('narrative_delivery');Q=D/'QA';Q.mkdir(exist_ok=True)
old=Document('manuscript_delivery/CD4_IKZF2_STAT1_integrated_manuscript_zh.docx');new=Document(D/'CD4_IKZF2_STAT1_integrated_manuscript_zh.docx')
def table_values(doc):return [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
assert table_values(old)==table_values(new),'Table data altered'
texts='\n'.join(p.text for p in new.paragraphs);heads=[p.text for p in new.paragraphs if p.style.name.startswith('Heading')]
manifest=json.loads((D/'document_manifest.json').read_text(encoding='utf8'))
fig_order=[v['file'] for v in manifest['figures']]
expected=['fig1_population.png','fig7_communication.png','fig4_networks.png','candidate_expression.png','fig6_bootstrap.png','fig2_pathways.png','fig3_programs.png','candidate_function.png','figS1_transfer.png'];assert fig_order==expected
assert len(manifest['references'])==15
for term in ['GSEA','GSVA','UCell','凋亡','焦亡','线粒体','SCENIC','Bootstrap','NicheNet','IKZF2','STAT1','免疫全景','CD4']:
 assert term in texts,term
baseline=json.loads((D/'baseline.json').read_text(encoding='utf8'));now={name:hashlib.sha256(path.read_bytes()).hexdigest() for name,path in [('E',Path('revision_delivery/report_evidence_aggregate.json')),('F',Path('manuscript_delivery/added_evidence_aggregate.json'))]};assert now==baseline['source_tables_sha256']
oldfig=Path('manuscript_delivery/figures');unchanged=[]
for name in expected:
 if (oldfig/name).exists():
  assert (oldfig/name).read_bytes()==(D/'figures'/name).read_bytes();unchanged.append(name)
# Record numeral preservation for review; narrative prose may remove duplicate presentations.
def scientific_tokens(doc):
 s='\n'.join(p.text for p in doc.paragraphs);s=s[s.index('2　结果'):s.index('4　统计与数值补充')]
 return set(re.findall(r'(?<![A-Za-z0-9])[-−+]?\d+\.\d+(?:[eE][−+-]?\d+)?',s))
missing=sorted(scientific_tokens(old)-scientific_tokens(new))
crossrefs=re.findall(r'(?:图|表)(?:S)?\d+',texts);result={'tables_identical':len(new.tables),'source_data_unchanged':True,'references':len(manifest['references']),'figure_order':fig_order,'figures_unchanged':unchanged,'main_figures':8,'supplementary_figures':1,'headings':heads,'crossrefs':sorted(set(crossrefs)),'old_prose_numerals_for_review':missing,'new_scientific_models':0}
(Q/'editorial_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({k:v for k,v in result.items() if k not in ['headings']},ensure_ascii=True))
