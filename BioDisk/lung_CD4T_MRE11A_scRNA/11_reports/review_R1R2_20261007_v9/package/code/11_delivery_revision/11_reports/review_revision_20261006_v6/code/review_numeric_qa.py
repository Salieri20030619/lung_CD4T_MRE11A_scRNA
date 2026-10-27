from pathlib import Path
from docx import Document
import pandas as pd, json,re,hashlib
B=Path(__file__).resolve().parent; D=B/'review_delivery'
S=Path('C:/Users/31937/tmp/pdfs/silicosis_audit_20261006/sources/11_reports')
doc=Document(D/'internal/revised_source.docx');checks=[]
def table(first):
 rows=[t for t in doc.tables if t.rows[0].cells[0].text==first];assert len(rows)==1
 return [[c.text for c in r.cells] for r in rows[0].rows[1:]]
def number(x):return float(x.replace('−','-'))
d=pd.read_csv(S/'human_first_20261004_v1/results/stage2_evidence_all_candidates.tsv.gz',sep='\t',low_memory=False)
d=d[(d.ligand=='IL15')&(d['filter']=='All')&(d.model=='Source')&(d.test=='Nominal05')&(d.direction=='Up')]
for row in table('接收端/来源'):
 rec,sen=row[0].split('/');z=d[(d.receiver==('Core' if rec=='Core' else 'CD4-IL7R'))&(d.sender=={'AM':'MacAM','MM':'MonoMac','MW':'MacWide'}[sen])].iloc[0]
 assert abs(number(row[1])-z.aupr_corrected)<0.00000051
 assert (row[2]=='通过')==bool(z.LR_expression_supported)
 assert (row[3]=='通过')==bool(z.all_chains_marginal_gate)
 if pd.notna(z.association_P):
  assert number(row[5])==z.association_n
  assert abs(number(row[6])-z.association_P)<0.00000051
  assert abs(number(row[7])-z.association_FDR)<0.00000051
 else:assert row[5:]==['—','未计算','未计算']
 checks.append('IL15_'+row[0])
d=pd.read_csv(S/'human_gap_diagnostics_20261005_v1/results/IL15_fixed_condition_attrition.tsv',sep='\t')
assert (d[(d.group=='CWP')&d.receiver.isin(['Core','CD4-IL7R'])&(d.stage=='FourMarginal10pct')].n==0).all()
d=pd.read_csv(S/'human_gap_diagnostics_20261005_v1/results/GZMA_IKZF2_source_composition_summary.tsv',sep='\t')
for row in table('来源/组'):
 source,group=row[0].split('/');z=d[(d.series==source)&(d.group==group)].iloc[0]
 assert int(row[1])==z.donors and row[2]==str(z.IKZF2_positive)+'/'+str(z.cells) and int(row[3])==z.IKZF2_umi
 checks.append('GZMA_'+row[0])
d=pd.read_csv(S/'technical_delivery_20261002_ikzf2_stat1_v1/revision_v2/final_release/data/S2_TF_associations.tsv',sep='\t')
d=d[(d.species=='Human')&(d.scope=='Core')&(d.tf=='IKZF2')&(d['filter']=='All')&(d.variant=='DropOverlap')&(d.model=='SourceComp')]
for row in table('程序'):
 z=d[(d.feature=='score_'+row[0])&(d.seed==int(row[1]))].iloc[0]
 for text,value in zip(row[2:5],[z.rho,z.p,z.FDR]):assert abs(number(text)-value)<0.00000051
 assert row[5]=='17/8' and z.n==17 and z.df==8
 checks.append('SourceComp_'+row[0]+'_'+row[1])
d=pd.read_csv(S/'human_gap_diagnostics_20261005_v1/results/IKZF2_matched_reference_contrast_models.tsv',sep='\t')
d=d[(d.variant=='DropOverlap')&(d.endpoint=='excess')].sort_values('seed')
assert d.n.eq(16).all() and d.df.eq(10).all()
assert [round(x,3) for x in d.rho]==[-.635,-.762,-.598]
assert [round(x,5) for x in d.P]==[.02656,.00398,.04014]
checks.append('excess_original_three_seeds')
text='\n'.join(p.text for p in doc.paragraphs)
fig_refs=set(re.findall(r'图(S?\d+)',text));tab_refs=set(re.findall(r'表(S?\d+)',text))
assert fig_refs<=set(map(str,range(1,11)))|{'S1','S2','S3','S4'}
assert tab_refs<=set(map(str,range(1,5)))|{'S'+str(x) for x in range(1,15)}
assert 'Library交付仅为报告PDF和可编辑文档' not in text
assert '终端效应程序为去共享基因后的 PRF1、IFNG、NKG7' in text
assert hashlib.sha256((B/'graphics_delivery/CD4_IKZF2_STAT1_integrated_manuscript_zh.docx').read_bytes()).hexdigest()=='1bd05f6e28a5c60da9a012b5397f237e0999919d43e0890b893074314882476c'
out={'checks':checks,'aggregate_rows_checked':19,'baseline_endpoint_checked':True,'cross_references_valid':True,'old_word_unchanged':True,'new_models':0}
(D/'QA/numeric_revision_checks.json').write_text(json.dumps(out,indent=2),encoding='utf-8')
print(json.dumps(out))
