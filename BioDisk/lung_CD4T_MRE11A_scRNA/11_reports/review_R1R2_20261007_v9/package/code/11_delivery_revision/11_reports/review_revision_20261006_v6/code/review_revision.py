from pathlib import Path
from copy import deepcopy
import json, hashlib, re, zipfile
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

B=Path(__file__).resolve().parent
D=B/'review_delivery'
for name in ['internal','QA','provenance']:(D/name).mkdir(parents=True,exist_ok=True)
source=B/'graphics_delivery/CD4_IKZF2_STAT1_integrated_manuscript_zh.docx'
old=Document(source); doc=Document(source)
baseline_hash=hashlib.sha256(source.read_bytes()).hexdigest()
changes=[]
def find(prefix):
    rows=[p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(rows)==1,(prefix,len(rows))
    return rows[0]
def put(prefix,text):
    p=find(prefix); before=p.text
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if props is not None:r._r.insert(0,props)
    changes.append({'before':before,'after':text})
    return p
def after(p,text,style='Normal'):
    e=OxmlElement('w:p');p._p.addnext(e);q=Paragraph(e,p._parent);q.style=style;q.add_run(text)
    q.paragraph_format.space_after=Pt(5)
    return q
def table_after(p,title,headers,rows,widths):
    cap=after(p,title,'Caption');cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,(c,text) in enumerate(zip(t.rows[0].cells,headers)):
        c.text=text;c.width=Cm(widths[i]); sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF4');c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        for i,(c,text) in enumerate(zip(t.add_row().cells,row)):c.text=str(text);c.width=Cm(widths[i])
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.first_line_indent=Pt(0)
                for r in pp.runs:r.font.size=Pt(8.3)
    cap._p.addnext(t._tbl)
    note=OxmlElement('w:p');t._tbl.addnext(note)
    return Paragraph(note,cap._parent)

put('尘肺患者 CD4 T 细胞转录状态','煤工尘肺患者与小鼠硅尘模型：\nCD4 T细胞状态、巨噬信号与IKZF2、STAT1关联')
put('人类患者数据与小鼠硅尘模型的整合分析','以人类患者为主的探索性整合分析')
put('数据及文献核对截至','数据及文献核对截至2026年10月6日｜审阅修订版')
put('目的：',find('目的：').text.replace('以尘肺患者','以煤工尘肺患者'))
put('方法：',find('方法：').text.replace('以供者为统计单位分析组成及程序差异','在来源强混杂的整合队列中，以供者为统计单位分析组成及程序差异'))
put('结果：',find('结果：').text.replace('巨噬信号分析中，IL15','接收端配体预测中，IL15').replace('；GZMA的FDR上调靶集则','，其巨噬来源与受体链证据分层见表2；GZMA的FDR上调靶集则'))
put('小鼠主要疾病对照来自',find('小鼠主要疾病对照来自').text.replace('每组 3 个独立文库','每组 3 份scRNA文库')+' 已核对六个命名文库及GSM条目；动物个体或混样与文库的对应关系尚未闭合，因此不将六文库写作已证实的六只独立动物。')
put('以供者或独立文库为推断单位',find('以供者或独立文库为推断单位').text.replace('以供者或独立文库为推断单位','人类以供者、小鼠以文库为分析单位'))
p=find('同供者关联另要求')
after(p,'配体活性使用各巨噬来源通过单对LR门槛的配体并集，在固定接收端疾病靶集及背景上计算一次；相同AUPR随后关联到各来源记录，不是发送端特异的三次活性计算。IL15的接收端复合物检查要求IL2RB及IL2RG各自在至少两名合格病例中达到10%检出门槛，不要求这两组病例相同。另保存严格四基因同供者条件：两端各≥20细胞，发送端IL15及IL15RA、接收端IL2RB及IL2RG的边际检出率均≥10%；它与单对LR、接收端链门槛及同细胞共检出是不同记录。')
p=find('独立时间点分析冻结')
after(p,'IL7R背景扣除疾病程序定义为：来源校正、名义P<0.05且log₂FC≥0.5的88基因上调靶集平均log1p(CP10k)评分，减去三组已固定匹配参照评分的等权均值；每组参照含88基因。该excess终点与IKZF2去重叠AUC在固定16名供者（9病例、7对照）中进行偏秩相关，调整疾病、来源及接收者秩深度。它是在同一数据内选择疾病程序后的特异性诊断，原基线仅有未校正探索性P，无FDR；三种子及原始程序、参照、excess和加入参照协变量的24条诊断记录均保留。10%/20%调控子成员删除沿用同一excess终点，六项预定高权重删除关联组成独立探索性BH家族；其余匹配删除对照为描述性。')
put('对两个已冻结的探索关系分别计算',find('对两个已冻结的探索关系分别计算').text+' 本节Bootstrap三基因终点与方法1.6、结果2.6的88基因背景扣除疾病程序分别定义，不互换解释。')
put('2.6　人类 IKZF2 的表达与 IL7R 亚群效应程序关联','2.6　人类IKZF2表达与IL7R背景扣除疾病程序的关联')
put('在IL7R亚群固定16人中','在IL7R亚群固定16人中，IKZF2去重叠调控子活性与背景扣除后的88基因疾病程序呈三个种子方向一致的负相关：ρ=−0.635/−0.762/−0.598，P=0.02656/0.00398/0.04014。该组为同数据诊断中的名义关联，原分析未计算FDR。预定移除10%/20%调控子成员后，与同一背景扣除终点的相关方向继续保留。表达和疾病程序关联分别定位于GZMA与IL7R，为IKZF2提供了两个明确的亚群观察。')
put('STAT1的主要人类线索来自',find('STAT1的主要人类线索来自').text.replace('与IKZF2的表达和效应程序关联','与IKZF2的表达和背景扣除疾病程序关联'))
put('IKZF2在人类具有GZMA RNA',find('IKZF2在人类具有GZMA RNA').text.replace('IL7R效应程序负相关','IL7R背景扣除疾病程序负相关'))
put('在人体核心CD4名义上调靶集中',find('在人体核心CD4名义上调靶集中').text) if any(p.text.startswith('在人体核心CD4名义上调靶集中') for p in doc.paragraphs) else None
put('在人类核心CD4名义上调靶集中',find('在人类核心CD4名义上调靶集中').text+' 这些AUPR是接收端共用预测值，逐来源表达与关联证据见表2及讨论3.2。')
put('凋亡通路的支持集中于特定调控条目',find('凋亡通路的支持集中于特定调控条目').text.replace('更宽的GO凋亡执行条目，','更宽的GO凋亡执行与KEGG凋亡亦呈名义负向富集（原始P均<0.05，见表S4）；GO凋亡执行条目，'))
put('完整档案包含：','完整档案包含原504项程序模型、人类通讯全部候选及靶集/背景/配体–靶成员、小鼠源表、12个网络完整成员、重采样诊断和GSEA/GSVA及敏感性结果。个体数据、大型对象留存服务器；本次Library仅更新PDF，下载目录另提供同一PDF及代码ZIP（含运行顺序与依赖说明），不交付Word。代码包需外部原数据和模型数据库，未进行完整流程重跑。')

p=find('IL15在部分人类名义靶集')
p=after(p,'IL15的预测优先级可保留，但其不同巨噬来源的证据并不相同。MacAM中仅一名合格病例的IL15达到10%检出，未通过至少两名病例的发送端门槛；因此其同供者关联未计算，不能写成阴性P。MonoMac和MacWide的IL15有两名病例达到门槛；IL7R中只有IL2RG支持单对LR，IL2RB达标病例为零。Core的IL2RB/IL2RG边际链门槛通过，仍未建立IL15RA转呈。严格四基因同供者10%条件在表2六种组合中均无合格病例；这是单独的覆盖诊断，不推断低检出受体在生物学上不存在。')
p=table_after(p,'表2　IL15接收端共用预测与逐巨噬来源证据',
 ['接收端/来源','校正AUPR','单对LR','β/γ链','四基因病例','关联n','P','原BH'],
 [['Core/AM','0.012640','未通过','通过','0','—','未计算','未计算'],['Core/MM','0.012640','通过','通过','0','14','0.620085','0.965810'],['Core/MW','0.012640','通过','通过','0','16','0.823641','0.982214'],['IL7R/AM','0.012361','未通过','未通过','0','—','未计算','未计算'],['IL7R/MM','0.012361','通过','未通过','0','14','0.101102','0.830072'],['IL7R/MW','0.012361','通过','未通过','0','15','0.294308','0.879680']],
 [2.7,2.5,1.6,1.6,1.7,1.2,2.9,2.9])
p.text='AM=MacAM，MM=MonoMac，MW=MacWide；MW为前两者的并集。β/γ链指接收端IL2RB/IL2RG边际门槛；四基因病例指严格同供者10%条件。均为完整细胞、来源校正、名义上调靶集；Core/IL7R靶数为126/88、背景为3739/2347，配体排名均为1。关联为配体表达与疾病程序的同供者偏秩相关，调整疾病、来源及两端深度；Core/IL7R的BH家族分别含253/232项有限P的配体×来源检验，不是本表六行另算BH。AUPR自身不作P/BH检验。'
p.style='Caption'

p=find('IKZF2的人类表达关联具有来源依赖性')
p=after(p,'来源覆盖进一步定位了这一表达信号（表3）。达到GZMA≥20细胞的实际模型名单为12人：G227的2名病例和7名对照，以及G288的3名病例；G135无人达到门槛。可辨识的来源内病例/对照交叉来自G227，来源校正效应应表述为探索性表达关联，不能推广为不同患者来源普遍上调。G288缺少同来源对照，零检出不支持该来源疾病下调。细胞计数用于描述检出覆盖，不作为生物学样本量。')
p=table_after(p,'表3　GZMA亚群IKZF2的来源检出与模型供者构成',
 ['来源/组','有细胞供者','检出/总细胞','总UMI','≥20细胞模型供者'],
 [['G227/CWP','2','16/369','17','2'],['G227/Control','8','13/1891','14','7'],['G288/CWP','6','0/292','0','3'],['G135/Control','1','0/2','0','0']], [3.2,3,4,2.5,4.4])
p.text='G135/G227/G288对应GSE135893/GSE227136/GSE288887；CWP为煤工尘肺，Control为对照。总计2554个GZMA细胞、29个IKZF2检出细胞、31 UMI；“有细胞供者”与模型纳入名单不同。';p.style='Caption'
p=find('背景筛选后有效供者降至15人')
p=after(p,'组成敏感性使用项目已保存的模型：在疾病、来源和秩深度基础上加入Core内部亚群比例（为避免共线性省略一个比例）。完整细胞、去重叠调控子仍为17名供者，残差df降至8（表4）。焦亡评分的三个相关仍负向，种子17/29保留名义P<0.05，均未通过各自原BH；凋亡评分三个种子不再达到名义阈值。Naive评分的正向关联亦有种子差异。这些结果支持保留程序关联及其组成敏感性，但不能把去基因重叠或一次组成调整解释为排除全部细胞状态贡献，更不能称为IKZF2抗凋亡、抗焦亡或因果保护作用。')
p=table_after(p,'表4　Core中IKZF2与程序的既有组成调整关联',
 ['程序','种子','ρ','P','原BH','n/df'],
 [['PyroExec','17','−0.800944','0.005360','0.175738','17/8'],['PyroExec','29','−0.779916','0.007794','0.375851','17/8'],['PyroExec','43','−0.519657','0.123682','0.599610','17/8'],['ApoExec','17','−0.419645','0.227318','0.646414','17/8'],['ApoExec','29','−0.451926','0.189779','0.657578','17/8'],['ApoExec','43','−0.438542','0.204869','0.664984','17/8'],['Naive','17','0.629543','0.051124','0.366928','17/8'],['Naive','29','0.582651','0.077128','0.524543','17/8'],['Naive','43','0.722516','0.018260','0.306630','17/8']], [3.3,1.5,3,3.5,3.5,2.3])
p.text='PyroExec/ApoExec/Naive为焦亡执行相关、凋亡执行相关及初始状态转录评分。保留原物种×背景×种子×调控子版本×范围×模型的完整TF–特征检验BH家族；未按本表九行重新校正。与图10及表S5的仅疾病/来源/深度主模型分开解释。';p.style='Caption'
p=find('小鼠核心CD4完整排序中')
after(p,'GSE314943的3+3为文库口径。现有元数据确认六个命名scRNA文库，但动物个体/混样到文库的映射未闭合；当前未证实伪重复，也不能据此把文库数直接写作独立动物数。该缺口限制生物学重复层面的确定性，跨物种结论仍定位于亚群和程序方向对照。')

# Preserve every original figure and original table; edits only concern documented paragraphs and three aggregate tables.
for r in doc.element.iter(qn('w:r')):
    pr=r.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
    fonts=pr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
    for k in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+k),'Times New Roman')
    fonts.set(qn('w:eastAsia'),'SimSun')
out=D/'internal/revised_source.docx';doc.save(out)
def matrices(d):return [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
om=matrices(old);nm=matrices(doc)
assert all(m in nm for m in om) and len(nm)==18
def media(p):
    with zipfile.ZipFile(p) as z:return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
assert media(source)==media(out) and len(doc.inline_shapes)==14
texts='\n'.join(p.text for p in doc.paragraphs)
for k in ['2.6　人类 IKZF2 的表达与 IL7R 亚群效应程序关联','IL7R效应程序负相关']:assert k not in texts
oldnums=set(re.findall(r'[-−+]?\d+\.\d+(?:[eE][−+-]?\d+)?','\n'.join(p.text for p in old.paragraphs)))
assert oldnums<=set(re.findall(r'[-−+]?\d+\.\d+(?:[eE][−+-]?\d+)?',texts))
assert hashlib.sha256(source.read_bytes()).hexdigest()==baseline_hash
result={'baseline_docx_sha256':baseline_hash,'changed_paragraphs':changes,'original_tables_preserved':15,'new_tables':3,'figures_preserved':14,'figure_bytes_identical':True,'original_decimal_values_preserved':True,'new_models':0,'internal_docx_not_for_delivery':str(out)}
(D/'provenance/revision_manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k!='changed_paragraphs'},ensure_ascii=True))
